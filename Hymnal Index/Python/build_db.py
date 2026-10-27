#!/usr/bin/env python3
"""Build hymnal.db (SQLite) from the hymn index spreadsheets.

Usage: python3 build_db.py

The spreadsheets stay the source of truth; re-run this script to regenerate
the database (and link_review.csv) after editing any of them.

Sources
  Lutheran Hymn Index.xlsx               the master index: one row per hymn, its other English
                                         titles in the 'Alternative English Title N' columns
  KELG-Hymn-Index.xlsx, 'Sheet 1'        KELG hymns with ELHB/TLH/LW/ELH/LSB counterparts
  ELHB Hymn Index.xlsx                   ELHB hymns by number and title
  wh_index.csv                           index of WH.pdf, made by extract_wh.py
  tlhy_index.csv                         TLHy's index of first lines (TLHy.pdf pp. 670-679)
  tlh_1941.csv                           TLH titles, authors, translators and original first
                                         lines (Project Wittenberg), made by extract_tlh.py
  tlhy_originals.csv                     what TLHy's translated hymns were translated from
                                         (researched; see its source and note columns)
  kelg_english.csv                       English translations found for KELG hymns that WH
                                         lacks (the 1917 supplement)
  te_laudamus_index.csv                  Te Laudamus's index of hymn titles (the photos in
                                         Te Laudamus/)
  tl_originals.csv                       what TL hymns that no title placed were translated
                                         from (researched, like tlhy_originals.csv)
  elgbs_lpc_hymns.csv                    the house hymnal of ELGBS and LPC's translation of it
                                         (the same 106 numbers): German and English first lines
  ws69_index.csv, ws2000_index.csv       the TLH supplements Worship Supplement (1969) and Worship
                                         Supplement 2000: number, first line, tune (Hymnary.org's
                                         tables WS69.html and WS2000.html, made by extract_ws.py)
  <CODE>/<CODE>_index.csv                the hymnals saved from Hymnary.org a page at a time
                                         (ACS/, AH/, ... TCSSH/): number, first line, tune (made by
                                         extract_hymnary.py); one is linked once its column is in
                                         the master index (write_back.py --add CODE puts it there)
  tch_index.csv                          The Concordia Hymnal's own index of first lines (the
                                         screenshots TCH/index1-8.jpg), used for TCH in place of
                                         Hymnary's table, which lacks the book's numbers
  lh_index.csv                           Liber Hymnorum's hymns (LH4A-C.pdf): English first line,
                                         Latin original, use, author, translator; read like a
                                         Hymnary.org table, its Latin matched as an original too
  link_decisions.csv                     human decisions on links: 'match' or 'no match'
  same_hymn_decisions.csv                which master-index rows are one hymn ('same hymn': write_back.py
                                         merges them), which only share numbers ('different hymns'), and
                                         which are one hymn kept on rows of their own ('same hymn, own rows')
  master_corrections.csv                 wrong numbers and misspelled titles in the master index
                                         (write_back.py corrects them; read here for the old titles)

Schema
  hymnals       one row per hymnal column in the master index
  hymns         one row per master-index row (id = spreadsheet row number), with its
                KELG or ELGBS German title, original title and language (for hymns
                that aren't KELG or ELGBS hymns), and other English title
  hymn_alternative_titles  the row's 'Alternative English Title N' cells
  hymn_cells    the raw cell text for each (hymn, hymnal) pair, e.g. '656-657'
  hymn_numbers  each cell split into individual numbers, e.g. 656 and 657
  hymn_index    view that pivots hymn_cells back into the spreadsheet layout (the
                alternative titles joined with ' | ')

  kelg_hymns    KELG/WH index: number, German title, WH English title, verse count
  kelg_refs     each KELG hymn's counterpart in ELHB/TLH/LW/ELH/LSB (number, title, verses)
  kelg_index    view: one row per KELG/WH hymn with its German title, WH title,
                its ELHB, TLH, LW, ELH and LSB numbers and titles, and any other
                English translation (kelg_english.csv)
  hymn_titles   every title we know for a master-index row: its HYMN TITLE and
                alternative titles plus KELG (German), WH, ELHB, TLHy, TLH, LW, ELH,
                LSB, TL, LPC, ELGBS (German, by LPC number) titles from the KELG, ELHB,
                TLHy, TL and ELGBS/LPC indexes, and the English first lines of the
                Hymnary.org hymnals (WS69, WS00, ACS, ...)
  hymn_titles_by_hymnal  view: hymn_titles with one column per source
  hymn_catalog  view: one row per master-index row with every hymnal's number,
                the German (KELG or ELGBS) title, the English title in the master
                index and its alternative titles, the title in each hymnal we have an
                index of, the original title and language, and the other English
                title (null where unknown)
  elhb_hymns    ELHB index: number, title
  tlhy_index    TLHy index of first lines: number, first line, translator
                ('alternate' marks a second index entry for the same hymn)
  tlhy_originals  tlhy_originals.csv: each researched TLHy hymn's original language,
                first line and author, with confidence and source
  tlhy_status   view: each TLHy hymn with the master-index rows it was matched to,
                or why it couldn't be (see link_review.csv)
  tlhy_sorting  each TLHy hymn sorted: a KELG hymn, translated from another
                language, an English original, or a translation not yet traced
  tl_index      Te Laudamus index of hymn titles: number ('157a' keeps its letter), title
                ('alternate' marks a cross-reference entry)
  tl_originals  tl_originals.csv, laid out like tlhy_originals
  tl_status     view: each TL hymn with the master-index rows it was matched to
  elgbs_lpc_hymns  elgbs_lpc_hymns.csv: ELGBS's Hausgesangbüchlein and LPC's translation,
                numbered alike (the master index has an LPC# column; ELGBS gives the
                German first lines)
  lpc_originals  view: the German first line of each LPC hymn, laid out like tlhy_originals
  hymnary_index the hymns of the hymnals taken from Hymnary.org's tables (hymnal 'WS69',
                'WS00', 'ACS', ...): number ('718a' keeps its letter), first line as
                Hymnary gives it, the English first line taken from it, the first line in
                another language where Hymnary gives both, tune ('alternate' marks another
                first line the book's own index gives a part of a hymn)
  hymnary_status  view: each of those hymns with the master-index rows it was matched to
  tlh_hymns     TLH (1941) from tlh_1941.csv: title, first line, author, translator,
                original first line (ASCII, umlauts as ae/oe/ue), tune
  kelg_english  kelg_english.csv: English translations of KELG hymns WH lacks
  wh_hymns      Walther's Hymnal (WH.pdf, via wh_index.csv): English and German
                first lines, tune, and the tune settings it points to
  wh_tune_settings  those tune settings, one row per TLH/LW/LSB number

  links         KELG/ELHB/TLHy/TL/LPC numbers and those of the Hymnary.org hymnals for
                master-index rows: those already in the master index plus proposed
                ones, with how each was found and whether it needs a human look ('ok' or
                'review'), or was decided against in link_decisions.csv ('rejected')
  proposed_index  hymn_index plus the linked KELG, ELHB, TLHy, TL, LPC, Hymnary-hymnal and
                German title columns
"""
import csv
import difflib
import re
import sqlite3
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
LCMS_XLSX = HERE / "Lutheran Hymn Index.xlsx"
KELG_XLSX = HERE / "KELG-Hymn-Index.xlsx"
ELHB_XLSX = HERE / "ELHB Hymn Index.xlsx"
WH_CSV = HERE / "wh_index.csv"  # extracted from WH.pdf by extract_wh.py
TLHY_CSV = HERE / "tlhy_index.csv"  # transcribed from TLHy.pdf, checked against two OCR passes
TLH_CSV = HERE / "tlh_1941.csv"  # TLH notes from Project Wittenberg, made by extract_tlh.py
TLHY_ORIGINALS_CSV = HERE / "tlhy_originals.csv"  # researched originals of TLHy's translations
KELG_ENGLISH_CSV = HERE / "kelg_english.csv"  # English versions of KELG hymns WH lacks
TL_CSV = HERE / "te_laudamus_index.csv"  # transcribed from the photos in Te Laudamus/, checked against OCR
TL_ORIGINALS_CSV = HERE / "tl_originals.csv"  # researched originals of TL hymns no title placed
ELGBS_LPC_CSV = HERE / "elgbs_lpc_hymns.csv"  # ELGBS's house hymnal and LPC's translation, numbered alike
# The indexes taken from Hymnary.org's tables, by hymnal code: the two TLH supplements (made
# by extract_ws.py), then the hymnals saved a page at a time (extract_hymnary.py), in the
# order they were added to the master index. One is linked once its column is there.
HYMNARY_CSVS = {"WS69": HERE / "ws69_index.csv", "WS00": HERE / "ws2000_index.csv"}
HYMNARY_CSVS.update({code: HERE / code / f"{code}_index.csv" for code in (
    "ACS", "AH", "ALH", "CSB", "ELHal", "ELW", "HCH", "HS91", "SBH", "THOS",
    "TBW07", "TBW67", "CHPPPW", "CHL", "TCSSH")})
# Hymnary has only placeholder numbers for TCH, so its numbers come from the book's own
# index of first lines, transcribed from TCH/index1-8.jpg and checked against Hymnary's list.
HYMNARY_CSVS["TCH"] = HERE / "tch_index.csv"
# Liber Hymnorum's English first lines and their Latin originals, from LH4A-C.pdf; its
# hymns are matched by both, and by the Latin through link_by_original (lh_originals).
HYMNARY_CSVS["LH"] = HERE / "lh_index.csv"
LINK_DECISIONS_CSV = HERE / "link_decisions.csv"  # human decisions on review items
SAME_HYMN_CSV = HERE / "same_hymn_decisions.csv"  # which master-index rows are one hymn
MASTER_CORRECTIONS_CSV = HERE / "master_corrections.csv"  # wrong numbers and titles in the master index
DB_PATH = HERE / "hymnal.db"
REVIEW_CSV = HERE / "link_review.csv"

# Master-index column header -> (code, full name, year)
HYMNALS = {
    "KELG#": ("KELG", "Kirchen-Gesangbuch für Evangelisch-Lutherische Gemeinden (Walther's Hymnal)", 1862),
    "TLHy#": ("TLHy", "The Lutheran Hymnary", 1913),
    "ELHB#": ("ELHB", "Evangelical Lutheran Hymn-Book", 1912),
    "TL#": ("TL", "Te Laudamus", 2024),
    "LPC#": ("LPC", "Lutheran Prayer Companion", 2018),
    "LSB#": ("LSB", "Lutheran Service Book", 2006),
    "TLH#": ("TLH", "The Lutheran Hymnal", 1941),
    "WS'69#": ("WS69", "Worship Supplement", 1969),
    "WS'00#": ("WS00", "Worship Supplement 2000", 2000),
    "LW#": ("LW", "Lutheran Worship", 1982),
    "HS'98#": ("HS98", "Hymnal Supplement 98", 1998),
    "ELH#": ("ELH", "Evangelical Lutheran Hymnary", 1996),
    "CW'21#": ("CW21", "Christian Worship: Hymnal", 2021),
    "CWS#": ("CWS", "Christian Worship: Supplement", 2008),
    "CW'93#": ("CW93", "Christian Worship: A Lutheran Hymnal", 1993),
    "LBW#": ("LBW", "Lutheran Book of Worship", 1978),
    "WOV#": ("WOV", "With One Voice", 1995),
    # from Hymnary.org's tables (extract_hymnary.py); the year is Hymnary's edition
    "ACS#": ("ACS", "All Creation Sings (a supplement to ELW)", 2020),
    "AH#": ("AH", "Ambassador Hymnal (AFLC)", 1994),
    "ALH#": ("ALH", "American Lutheran Hymnal", 1930),
    "CSB#": ("CSB", "Common Service Book", 1918),
    "ELHal#": ("ELHal", "Evangelical Lutheran Hymnal", 1908),
    "ELW#": ("ELW", "Evangelical Lutheran Worship", 2006),
    "HCH#": ("HCH", "Hymnal for Church and Home (4th ed., expanded)", 1938),
    "HS'91#": ("HS91", "Hymnal Supplement 1991 (a second supplement to LBW)", 1991),
    "SBH#": ("SBH", "Service Book and Hymnal", 1958),
    "THOS#": ("THOS", "The Hymnal and Order of Service", 1937),
    "TBW'07#": ("TBW07", "The Book of Worship", 1907),
    "TBW'67#": ("TBW67", "The Book of Worship", 1867),
    "CHPPPW#": ("CHPPPW", "A Collection of Hymns and Prayers, for Public and Private Worship", 1845),
    "CHL#": ("CHL", "A Collection of Hymns and a Liturgy", 1827),
    "TCSSH#": ("TCSSH", "The Church and Sunday-School Hymnal", 1898),
    "TCH#": ("TCH", "The Concordia Hymnal (revised)", 1960),
    "LH#": ("LH", "Liber Hymnorum: The Latin Hymns of the Lutheran Church", 2016),
}

# KELG 'Sheet 1' layout: each hymnal is a (number, title, verses) column triple.
KELG_SHEET = "Sheet 1"
KELG_REF_COLUMNS = {"ELHB": 3, "TLH": 6, "LW": 9, "ELH": 12, "LSB": 15}
KELG_WH_TITLE_COLUMN = 18  # S: English first line in Walther's Hymnal
KELG_WH_TITLE_HEADER = "Walther's Hymnal (2012)"
# Hymnals whose numbers are used to find a KELG hymn's rows in the master index.
# ELHB is left out: the master index barely has ELHB numbers yet.
KELG_MATCH_HYMNALS = ("TLH", "LW", "ELH", "LSB")

# Hymnals this script links to master-index rows (the rest are the master's own).
LINKED_HYMNALS = ("KELG", "ELHB", "TLHy", "TL", "LPC", *HYMNARY_CSVS)
# Hymnals that number separate items 157a and 157b: their links keep the letter.
# (CSB, ELHal, HS91 and SBH print a text set to two tunes twice, as 20a and 20b.)
LETTERED_HYMNALS = ("TL", "WS69", "CSB", "ELHal", "HS91", "SBH")

# Numbers the ELHB index gets wrong: number -> title. Checked against the ELHB
# itself (1927 printing, https://archive.org/details/evangelicalluthe0000vari_b7v6:
# hymn headings and author index); the KELG index agrees.
ELHB_INDEX_FIXES = {
    230: "Lord Jesus Christ, strong Hero Thou",  # the index swaps 230 and 286
    286: "Lord God, we all to Thee give praise",
    431: "May God be praised henceforth and blest forever",  # and 431 and 461
    461: "Lord, in Thy kingdom there shall be",
}

# Scanning errors and cut-off first lines in the ELHB index, fixed in the
# titles taken from it (hymn_titles, and the rows write_back.py adds for it).
ELHB_TITLE_FIXES = {
    44: "God, who madest earth and heaven, Darkness and light",
    201: "O bleeding Head, and wounded",
    207: "Seven times our blessed Savior spoke",
    300: "O bless, Thou heavenly Potentate",
    324: "Nearer, my God, to Thee, through Word and Sacrament",
    572: "My soul doth magnify the Lord",
    585: "Thee, Eternal God most high",
}

# Typing errors in Hymnary.org's first lines, fixed in hymnary_index (so in the titles
# taken from them): (hymnal, number) -> (the first line as Hymnary gives it, fixed).
# CHL's lines that run on into the second line are cut back to the first line, and a
# few lines Hymnary cuts short are completed (CHL 19 and 118, TCSSH 48), so the rows and
# alternative titles made from them read cleanly.
HYMNARY_FIXES = {
    ("ACS", "914"): ('Hark! The herald andgels sing', 'Hark! The herald angels sing'),
    ("ACS", "939"): ('Touch that soothes and helas the hurting', 'Touch that soothes and heals the hurting'),
    ("ACS", "943"): ('As the wind song though the trees', 'As the wind song through the trees'),
    ("ACS", "949"): ('Keep you lamps trimmed and burning', 'Keep your lamps trimmed and burning'),
    ("ACS", "959"): ('Life-gving water soaking our dryness', 'Life-giving water soaking our dryness'),
    ("ACS", "966"): ('Loaves were broken, words wer spoken', 'Loaves were broken, words were spoken'),
    ("ACS", "990"): ('Ho, ho ho! Sizohamba naye! (Ho, ho ho! We will gow with God!)',
                     'Ho, ho ho! Sizohamba naye! (Ho, ho ho! We will go with God!)'),
    ("ACS", "996"): ('Tend your ailing ones; in yor love, Lord', 'Tend your ailing ones; in your love, Lord'),
    ("ACS", "1000"): ('Gods work, our hands', "God's work, our hands"),
    ("ACS", "1036"): ("Commonwelath is God's commandment", "Commonwealth is God's commandment"),
    ("ACS", "1038"): ('God, we gather as your people to raise our sing above',
                      'God, we gather as your people to raise our song above'),
    ("ACS", "1040"): ('Love has broght us here together', 'Love has brought us here together'),
    ("ACS", "1049"): ('Before the watrs nourished earth', 'Before the waters nourished earth'),
    ("ACS", "1060"): ('Gentle Josph heard a warning', 'Gentle Joseph heard a warning'),
    ("ACS", "1074"): ('Óyenos, Seńor', 'Óyenos, Señor'),
    ("ACS", "1080"): ('Lord, I will lift mine eyse to the hills', 'Lord, I will lift mine eyes to the hills'),
    ("ELHal", "24"): ('The Bridgegroom soon will call us', 'The Bridegroom soon will call us'),
    ("ELHal", "79"): ('Halelujah! lo, He wakes!', 'Hallelujah! lo, He wakes!'),
    ("ELHal", "283"): ('Bridgroom, Thou art mine', 'Bridegroom, Thou art mine'),
    ("ELHal", "325"): ('Come, O Creator, Sprit blest!', 'Come, O Creator, Spirit blest!'),
    ("ELHal", "355"): ("By cool Siloan's shady rill", "By cool Siloam's shady rill"),
    ("ELHal", "464"): ("On Jordan's stomy banks I stand", "On Jordan's stormy banks I stand"),
    ("ELW", "518"): ('Wisdom calls throught the city', 'Wisdom calls through the city'),
    ("ELW", "528"): ('Confitemini Domino (Come and fill our hearts with you peace)',
                     'Confitemini Domino (Come and fill our hearts with your peace)'),
    ("ELW", "534"): ('Savior, again to your dear name we rasie', 'Savior, again to your dear name we raise'),
    ("ELW", "706"): ('Los pobres siempre esperan el amanecer (The poor ones of the world awain the dawn of hope)',
                     'Los pobres siempre esperan el amanecer (The poor ones of the world await the dawn of hope)'),
    ("ELW", "823"): ("Praise the Lord! O heavn's adore him", "Praise the Lord! O heav'ns adore him"),
    ("HCH", "277"): ("Come, ye disconsolate, whre'er ye languish", "Come, ye disconsolate, where'er ye languish"),
    ("TBW07", "1"): ('We praise The, O God', 'We praise Thee, O God'),
    ("TBW07", "8"): ("I'll praise my Make while I've breath", "I'll praise my Maker while I've breath"),
    ("TBW07", "35"): ('This is the day the Lord hath nade', 'This is the day the Lord hath made'),
    ("TBW07", "358"): ('In holy comtemplation', 'In holy contemplation'),
    ("TBW07", "490"): ('Once was heard the song of childern', 'Once was heard the song of children'),
    ("TBW67", "36"): ('Soveriegn Ruler of the skies', 'Sovereign Ruler of the skies'),
    ("TBW67", "39"): ("Not from relentless fate's dakr womb", "Not from relentless fate's dark womb"),
    ("TBW67", "40"): ('Praise ye Jehovah! with anthems of praise come beofer Him',
                      'Praise ye Jehovah! with anthems of praise come before Him'),
    ("TBW67", "44"): ('The Lord is great! ye hosts of heavn, adore Him', "The Lord is great! ye hosts of heav'n, adore Him"),
    ("TBW67", "123"): ("On the mountain's to appearing", "On the mountain's top appearing"),
    ("TBW67", "257"): ('Sin, like a venmous disease', 'Sin, like a venomous disease'),
    ("TBW67", "321"): ('Abide with me! fast falls the evenitde', 'Abide with me! fast falls the eventide'),
    ("TBW67", "461"): ('When shall we all meet agin?', 'When shall we all meet again?'),
    ("TBW67", "157"): ('Twas on that dreadful, doleful night', "'Twas on that dreadful, doleful night"),
    ("TBW67", "289"): ('Sinner. is thy heart at rest?', 'Sinner, is thy heart at rest?'),
    ("TBW67", "334"): ('Tis my happiness below', "'Tis my happiness below"),
    ("TBW67", "373"): ('Mid scenes of confusion and creature complaints', "'Mid scenes of confusion and creature complaints"),
    ("CHPPPW", "110"): ('Self rightous souls on works rely', 'Self-righteous souls on works rely'),
    ("CHPPPW", "165"): ('How short and ahsty is our life!', 'How short and hasty is our life!'),
    ("CHPPPW", "315"): ('Hark! what trimphant strains are these', 'Hark! what triumphant strains are these'),
    ("CHPPPW", "201"): ('Tis by the faith of joys to come', "'Tis by the faith of joys to come"),
    ("CHL", "83"): ("Not from relentless fate's dar", "Not from relentless fate's dark womb"),  # cut off
    ("CHL", "191"): ('Are not thy mercies sovreign still', "Are not thy mercies sov'reign still"),
    ("CHL", "541"): ('Glory to God who reigns aove', 'Glory to God who reigns above'),
    ("CHL", "101"): ('praise ye the lord, who reigns', 'Praise ye the Lord, who reigns'),
    ("CHL", "142"): ('ôTis finished, so the Savior cried, and meekly', "'Tis finished, so the Savior cried, and meekly"),
    ("CHL", "292"): ('Ye servant [servants] of the Lord, Each in his [your] office waits', 'Ye servants of the Lord'),
    ("CHL", "413"): ('How rich thy gifts, Almighty King1', 'How rich thy gifts, Almighty King'),
    ("CHL", "695"): ("Heav'n is a place or rest from sin", "Heav'n is a place of rest from sin"),
    ("CHL", "19"): ('My soul, praise the Lord, speak good [well] of', 'My soul, praise the Lord, speak good of his name'),
    ("CHL", "33"): ('All powerful, self existent God, Who all creation', 'All powerful, self existent God'),
    ("CHL", "45"): ("Praise ye the Lord, 'tis good to raise your [our] hearts and voices", "Praise ye the Lord, 'tis good to raise"),
    ("CHL", "50"): ('The promises I [we] sing which love supreme [sovereign love] once [hath] spoke', 'The promises I sing'),
    ("CHL", "73"): ('Ye subjects of the Lord, proclaim The royal honors of his name', 'Ye subjects of the Lord, proclaim'),
    ("CHL", "96"): ("On Judah's plains as shepherds sat Watching [kept Watch o'er] their flocks",
                    "On Judah's plains as shepherds sat"),
    ("CHL", "103"): ('Sons of Adam, join to raise Songs of gratitude and praise', 'Sons of Adam, join to raise'),
    ("CHL", "118"): ('All-glorious God, what hymns of', 'All-glorious God, what hymns of praise'),
    ("CHL", "132"): ('God of mercy, God of love [grace], Hear', 'God of mercy, God of love'),
    ("CHL", "138"): ('The Savior, what a noble flame was kindled', 'The Savior, what a noble flame'),
    ("CHL", "139"): ('Behold the amazing sight, The Savior lifted high', 'Behold the amazing sight'),
    ("CHL", "140"): ('Come let us lift our voices high, High as our joys arise', 'Come let us lift our voices high'),
    ("CHL", "181"): ('Great God, arise and shine, with beams', 'Great God, arise and shine'),
    ("CHL", "189"): ('My hope my portion and my God how little art', 'My hope, my portion, and my God'),
    ("CHL", "218"): ('My God, my portion, and my love, my everlasting all.', 'My God, my portion, and my love'),
    ("CHL", "258"): ('Come, ye [you] [we] that [who] love the Lord [Christ], And let your [our] joys',
                     'Come, ye that love the Lord'),
    ("CHL", "330"): ('If solid happiness we prize, Within our breast', 'If solid happiness we prize'),
    ("CHL", "353"): ('Parent of all, omnipotent In heaven and earth', 'Parent of all, omnipotent'),
    ("CHL", "361"): ('Come let our voices join to raise a sacred song of solemn praise', 'Come let our voices join to raise'),
    ("CHL", "371"): ('Praise to God, the great Creator, Bounteous', 'Praise to God, the great Creator'),
    ("CHL", "406"): ("Indulgent God, whose bounteous care O'er all thy works is shown", 'Indulgent God, whose bounteous care'),
    ("CHL", "419"): ('O Lord, our fathers oft have told, In our attentive ears', 'O Lord, our fathers oft have told'),
    ("CHL", "434"): ('Consider all my sorrows, Lord, And thy deliverance send', 'Consider all my sorrows, Lord'),
    ("CHL", "454"): ('My God, my [mine] everlasting hope, I live upon thy truth', 'My God, my everlasting hope'),
    ("TCSSH", "69"): ('Come, thou longexpected Jesus', 'Come, thou long-expected Jesus'),
    ("TCSSH", "88"): ("Hail the Lord's Anointed", "Hail to the Lord's Anointed"),
    ("TCSSH", "48"): ('Heaven and earth, and air', 'Heaven and earth, and sea, and air'),
    ("TCH", "17"): ('Praise ye the Father', 'Praise ye the Father for His loving kindness'),  # its index cuts it short
}

# Master-index column holding each row's German title(s): KELG's, or for a hymn of
# ELGBS's house hymnal (a supplement to KELG in all but name) that KELG lacks, ELGBS's.
GERMAN_TITLE_HEADER = "KELG/ELGBS German Title"
# Text columns after the hymnal numbers in the master index -> hymns table column.
TEXT_COLUMNS = {
    GERMAN_TITLE_HEADER: "german_title",
    "Original Title": "original_title",        # when the hymn isn't a KELG or ELGBS hymn
    "Original Language": "original_language",
    "Other English Title": "other_english_title",  # a translation from outside these hymnals
    "Other English Source": "other_english_source",
}
# The master index's other English titles for a row: 'Alternative English Title 1', '... 2', ...
ALT_TITLE_HEADER = "Alternative English Title"
ALT_TITLE_RE = re.compile(rf"^{ALT_TITLE_HEADER} (\d+)$")
# master_corrections.csv 'hymnal' value for a correction of the HYMN TITLE itself
TITLE_CORRECTION = "title"
# master_corrections.csv 'hymnal' value for an Alternative English Title a row lacks
ALT_CORRECTION = "alternate"
# master_corrections.csv 'hymnal' value for a correction of the KELG/ELGBS German Title
GERMAN_CORRECTION = "german"
# tlhy_originals.csv confidence levels good enough to act on
TRUSTED = ("confirmed", "probable")

FUZZY_THRESHOLD = 0.85
# A title that is the start of more rows' titles than this is too common a start to
# link for review ('Lord Jesus Christ'; link_by_title).
PREFIX_LIMIT = 3

SCHEMA = """
CREATE TABLE hymnals (
    code          TEXT PRIMARY KEY,
    column_header TEXT NOT NULL,      -- header as it appears in the spreadsheet
    name          TEXT NOT NULL,
    year          INTEGER,
    sort_order    INTEGER NOT NULL    -- spreadsheet column order
);

CREATE TABLE hymns (
    id           INTEGER PRIMARY KEY, -- spreadsheet row number
    title        TEXT NOT NULL,
    scripture    TEXT,                -- the spreadsheet's 'Text' column
    german_title TEXT,                -- the 'KELG/ELGBS German Title' column
    original_title       TEXT,        -- 'Original Title': first line of the original, when
                                      -- the hymn isn't a KELG or ELGBS hymn
    original_language    TEXT,        -- 'Original Language'
    other_english_title  TEXT,        -- 'Other English Title': a translation found elsewhere
    other_english_source TEXT         -- 'Other English Source': its translator and source
);

CREATE TABLE hymn_alternative_titles (  -- the 'Alternative English Title N' columns
    hymn_id  INTEGER NOT NULL REFERENCES hymns(id),
    position INTEGER NOT NULL,        -- N
    title    TEXT NOT NULL,
    PRIMARY KEY (hymn_id, position)
);

CREATE TABLE hymn_cells (
    hymn_id     INTEGER NOT NULL REFERENCES hymns(id),
    hymnal_code TEXT NOT NULL REFERENCES hymnals(code),
    raw         TEXT NOT NULL,        -- cell as written, e.g. '656-657', 'p. 7,17'
    PRIMARY KEY (hymn_id, hymnal_code)
);

CREATE TABLE hymn_numbers (
    hymn_id     INTEGER NOT NULL,
    hymnal_code TEXT NOT NULL,
    position    INTEGER NOT NULL,     -- order within the cell, starting at 1
    number      INTEGER NOT NULL,
    suffix      TEXT NOT NULL DEFAULT '',  -- e.g. 'a' in '982a'
    is_page     INTEGER NOT NULL DEFAULT 0, -- 1 when the cell gives page numbers ('p. 29')
    PRIMARY KEY (hymn_id, hymnal_code, position),
    FOREIGN KEY (hymn_id, hymnal_code) REFERENCES hymn_cells(hymn_id, hymnal_code)
);

CREATE INDEX hymn_numbers_lookup ON hymn_numbers(hymnal_code, number);
CREATE INDEX hymns_title ON hymns(title COLLATE NOCASE);

CREATE TABLE kelg_hymns (
    number       INTEGER PRIMARY KEY,
    german_title TEXT NOT NULL,
    wh_title     TEXT,                -- English first line in Walther's Hymnal (WH)
    verses       TEXT,
    supplement   INTEGER NOT NULL DEFAULT 0  -- 1 for the hymns added in 1917
);

CREATE TABLE kelg_refs (
    kelg_number INTEGER NOT NULL REFERENCES kelg_hymns(number),
    hymnal_code TEXT NOT NULL,
    raw         TEXT NOT NULL,        -- number cell as written, e.g. '569 + 535'
    title       TEXT,
    verses      TEXT,
    fixed       TEXT,                 -- corrected number cell, when raw is a known mistake
    PRIMARY KEY (kelg_number, hymnal_code)
);

CREATE TABLE wh_hymns (              -- Walther's Hymnal (WH), numbered like KELG
    number             INTEGER PRIMARY KEY,
    english_first_line TEXT,
    german_first_line  TEXT,
    tune               TEXT,
    tune_settings      TEXT,         -- as printed, e.g. 'TLH 73, first tune; or LW 23*'
    first_stanza_line  TEXT,
    see_number         INTEGER       -- printed only as '(See # n.)'
);

CREATE TABLE wh_tune_settings (       -- tune_settings split up: where to find the tune
    number      INTEGER NOT NULL REFERENCES wh_hymns(number),
    hymnal_code TEXT NOT NULL,
    setting     INTEGER NOT NULL,
    alternate   INTEGER NOT NULL      -- 1 for '*': a substitute tune, not the appointed one
);

CREATE TABLE tlhy_index (            -- The Lutheran Hymnary (1913), index of first lines
    number      INTEGER NOT NULL,
    first_line  TEXT NOT NULL,
    translator  TEXT,
    alternate   INTEGER NOT NULL DEFAULT 0,  -- 1: a second index entry for the same hymn
    index_page  INTEGER,              -- printed page of the index
    note        TEXT
);

CREATE TABLE tlhy_originals (        -- tlhy_originals.csv
    number         INTEGER PRIMARY KEY,  -- TLHy number
    language       TEXT,
    original_title TEXT,              -- first line of the original
    author         TEXT,
    confidence     TEXT NOT NULL,     -- 'confirmed', 'probable', or 'partial' (no first line)
    same_as        TEXT,              -- the hymnal number of the very same translation,
                                      -- e.g. 'ELHB 456', where the research found one
    source         TEXT,
    note           TEXT
);

CREATE TABLE tlhy_sorting (          -- each TLHy hymn sorted by where it comes from
    number         INTEGER PRIMARY KEY,
    first_line     TEXT NOT NULL,
    translator     TEXT,
    status         TEXT NOT NULL,     -- 'matched', 'needs review' or 'unmatched' (tlhy_status)
    category       TEXT NOT NULL,     -- 'KELG hymn', 'translated from <language>',
                                      -- 'English original', 'translation, original not traced'
    kelg           TEXT,              -- KELG number(s), for KELG hymns
    original_title TEXT,
    original_language TEXT,
    original_author   TEXT
);

CREATE TABLE tl_index (              -- Te Laudamus (2024), index of hymn titles
    number      INTEGER NOT NULL,     -- a lettered number keeps its letter, as text: '157a'
    title       TEXT NOT NULL,
    alternate   INTEGER NOT NULL DEFAULT 0,  -- 1: a cross-reference entry ('Nicene Creed')
    index_page  INTEGER,              -- printed page of the index
    note        TEXT
);

CREATE TABLE tl_originals (          -- tl_originals.csv, laid out like tlhy_originals
    number         INT PRIMARY KEY,   -- TL number ('157a' as text; an INTEGER PRIMARY KEY
                                      -- would take only whole numbers)
    language       TEXT,
    original_title TEXT,
    author         TEXT,
    confidence     TEXT NOT NULL,
    same_as        TEXT,
    source         TEXT,
    note           TEXT
);

CREATE TABLE elgbs_lpc_hymns (       -- the Hausgesangbüchlein of ELGBS (15th ed., 1881) and LPC's
                                     -- translation of it (2018), which keeps its numbers
    number             INTEGER PRIMARY KEY,
    german_first_line  TEXT NOT NULL,   -- ELGBS, in LPC's spelling
    english_first_line TEXT NOT NULL,   -- LPC
    author             TEXT,
    translator         TEXT,            -- LPC's credit
    tune               TEXT,            -- LPC's tune line
    elgbs_page         INTEGER,         -- page in ELGBS's hymn appendix
    lpc_page           INTEGER,
    row_title          TEXT,            -- title of a new master row, where the first line would mislead
    note               TEXT
);

CREATE VIEW lpc_originals AS         -- laid out like tlhy_originals, for link_by_original
SELECT number, 'German' AS language, german_first_line AS original_title, author,
       'confirmed' AS confidence, NULL AS same_as, 'ELGBS' AS source, note
FROM elgbs_lpc_hymns;

CREATE TABLE hymnary_index (         -- the hymnals taken from Hymnary.org's tables: the TLH
                                     -- supplements (WS69, WS00) and the ones in HYMNARY_CSVS
    hymnal      TEXT NOT NULL,        -- 'WS69', 'WS00', 'ACS', ...
    number      INTEGER NOT NULL,     -- a lettered number keeps its letter, as text: '718a'
    first_line  TEXT NOT NULL,        -- as Hymnary gives it (typing errors fixed: HYMNARY_FIXES)
    english     TEXT NOT NULL,        -- the English first line in it (hymnary_parts)
    original    TEXT,                 -- the first line in another language, where it has both
    tune        TEXT,
    alternate   INTEGER NOT NULL DEFAULT 0,  -- 1: another first line the book's index gives it
    note        TEXT
);

CREATE VIEW lh_originals AS          -- Liber Hymnorum's Latin, laid out like tlhy_originals
SELECT number, 'Latin' AS language, original AS original_title, NULL AS author,
       'confirmed' AS confidence, NULL AS same_as, 'LH' AS source, note
FROM hymnary_index WHERE hymnal = 'LH' AND alternate = 0;

CREATE TABLE tlh_hymns (             -- The Lutheran Hymnal (1941), from tlh_1941.csv
    number         INTEGER PRIMARY KEY,
    title          TEXT NOT NULL,
    first_line     TEXT,
    author         TEXT,
    translator     TEXT,
    original_title TEXT,              -- original first line, ASCII as in the source (ae/oe/ue)
    tune           TEXT
);

CREATE TABLE kelg_english (          -- kelg_english.csv
    kelg_number INTEGER NOT NULL REFERENCES kelg_hymns(number),
    title       TEXT NOT NULL,
    translator  TEXT,
    source      TEXT NOT NULL,        -- where it was found, e.g. 'Hymnoglypt', 'Hymnary'
    source_url  TEXT,
    note        TEXT
);

CREATE TABLE elhb_hymns (
    number INTEGER PRIMARY KEY,
    title  TEXT NOT NULL
);

CREATE TABLE hymn_titles (           -- a master-index row's title in each hymnal we know it for
    hymn_id INTEGER NOT NULL REFERENCES hymns(id),
    source  TEXT NOT NULL,            -- 'master' (the HYMN TITLE column), 'alternate' (an Alternative
                                      -- English Title), 'KELG', 'WH', or a hymnal code
    title   TEXT NOT NULL,
    PRIMARY KEY (hymn_id, source, title)
);

CREATE TABLE links (
    hymn_id     INTEGER NOT NULL REFERENCES hymns(id),
    hymnal_code TEXT NOT NULL,        -- 'KELG', 'ELHB', 'TLHy', 'TL', 'LPC' or a Hymnary.org hymnal's ('WS69', ...)
    number      INTEGER NOT NULL,     -- a lettered number as text: '157a' (see link_number)
    method      TEXT NOT NULL,        -- how the link was found
    status      TEXT NOT NULL,        -- 'ok', 'review', or 'rejected' (link_decisions.csv)
    detail      TEXT,
    PRIMARY KEY (hymn_id, hymnal_code, number)
);
"""

NUMBER_RE = re.compile(r"^(\d+)([a-z]?)$")


def parse_cell(raw):
    """Split a hymn-number cell into [(number, suffix, is_page)].

    Handles '123', '656-657', '982-982a', '649, 975', '361/362', '569 + 535',
    '333\\n336', 'p. 7,17', 'p 198'. Ranges in these sheets only ever span two
    adjacent numbers, so a range yields its two endpoints. A trailing '?'
    (uncertain number) is dropped.
    """
    text = raw.strip().rstrip("?").strip()
    is_page = bool(re.match(r"p\.?\s*\d", text, re.I))
    if is_page:
        text = re.sub(r"^p\.?", "", text, flags=re.I)
    parts = [p.strip() for p in re.split(r"[,/+\-\n]", text) if p.strip()]
    out = []
    for part in parts:
        m = NUMBER_RE.match(part)
        if not m:
            raise ValueError(f"unparseable hymn number {raw!r}")
        out.append((int(m.group(1)), m.group(2), int(is_page)))
    return out


def link_number(number, suffix=""):
    """A hymn number as the links table keeps it: 241, or '157a' (text) when it
    has a letter. Te Laudamus prints two items under one number as 157a, 157b."""
    return f"{number}{suffix}" if suffix else number


def parse_number(text):
    """'241' -> 241, '157a' -> '157a'."""
    m = NUMBER_RE.match(text.strip())
    if not m:
        raise ValueError(f"unparseable hymn number {text!r}")
    return link_number(int(m.group(1)), m.group(2))


def number_order(n):
    """Sort key for link numbers, lettered or not."""
    m = re.match(r"(\d+)(.*)", str(n))
    return int(m.group(1)), m.group(2)


def cell_text(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def normalize_title(title):
    """Lowercase, strip accents and punctuation, collapse whitespace."""
    text = unicodedata.normalize("NFKD", title)
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = re.sub(r"[^\w\s]", " ", text.lower())
    return " ".join(text.split())


# --- loading ---------------------------------------------------------------

def load_lcms(con):
    ws = openpyxl.load_workbook(LCMS_XLSX, read_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [h.strip() if isinstance(h, str) else h for h in next(rows)]

    if header[0] != "HYMN TITLE" or header[1] != "Text":
        raise SystemExit(f"unexpected header: {header[:2]}")
    hymnal_cols = []
    text_cols = {}  # hymns column -> spreadsheet column index
    alt_cols = []  # (spreadsheet column index, N of 'Alternative English Title N')
    for idx, h in enumerate(header[2:], start=2):
        if h is None:
            continue
        if h in TEXT_COLUMNS:
            text_cols[TEXT_COLUMNS[h]] = idx
            continue
        m = ALT_TITLE_RE.match(str(h))
        if m:
            alt_cols.append((idx, int(m.group(1))))
            continue
        if h not in HYMNALS:
            raise SystemExit(f"unknown hymnal column {h!r}; add it to HYMNALS or TEXT_COLUMNS")
        hymnal_cols.append((idx, h))

    for order, (_, h) in enumerate(hymnal_cols, start=1):
        code, name, year = HYMNALS[h]
        con.execute("INSERT INTO hymnals VALUES (?,?,?,?,?)", (code, h, name, year, order))

    for row_num, row in enumerate(rows, start=2):
        if row[0] is None:
            if any(v is not None for v in row):
                raise SystemExit(f"row {row_num} has data but no title")
            continue
        title = str(row[0]).strip()
        texts = [cell_text(row[text_cols[c]]) if c in text_cols and text_cols[c] < len(row) else None
                 for c in TEXT_COLUMNS.values()]
        con.execute("INSERT INTO hymns VALUES (?,?,?,?,?,?,?,?)", (row_num, title, cell_text(row[1]), *texts))
        for idx, n in alt_cols:
            alt = cell_text(row[idx] if idx < len(row) else None)
            if alt:
                con.execute("INSERT INTO hymn_alternative_titles VALUES (?,?,?)", (row_num, n, alt))
        for idx, h in hymnal_cols:
            raw = cell_text(row[idx] if idx < len(row) else None)
            if raw is None:
                continue
            code = HYMNALS[h][0]
            con.execute("INSERT INTO hymn_cells VALUES (?,?,?)", (row_num, code, raw))
            for pos, (num, suffix, is_page) in enumerate(parse_cell(raw), start=1):
                con.execute(
                    "INSERT INTO hymn_numbers VALUES (?,?,?,?,?,?)",
                    (row_num, code, pos, num, suffix, is_page),
                )

    pivot = ",\n    ".join(
        f"MAX(CASE WHEN c.hymnal_code = '{HYMNALS[h][0]}' THEN c.raw END) AS \"{HYMNALS[h][0]}\""
        for _, h in hymnal_cols
    )
    con.execute(
        f"""CREATE VIEW hymn_index AS
SELECT h.id, h.title,
    (SELECT group_concat(a.title, ' | ') FROM
        (SELECT title FROM hymn_alternative_titles WHERE hymn_id = h.id ORDER BY position) a)
        AS alternative_titles,
    h.scripture, h.german_title,
    {pivot},
    h.original_title, h.original_language, h.other_english_title, h.other_english_source
FROM hymns h LEFT JOIN hymn_cells c ON c.hymn_id = h.id
GROUP BY h.id
ORDER BY h.id"""
    )


def load_kelg(con):
    ws = openpyxl.load_workbook(KELG_XLSX, read_only=True)[KELG_SHEET]
    supplement = 0
    for row_num, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        if row[0] is None:
            if cell_text(row[1]) and "1917" in row[1]:
                supplement = 1  # 'Hymns added in 1917:' divider
            elif any(cell_text(v) for v in row):
                raise SystemExit(f"KELG row {row_num} has data but no number")
            continue
        wh_title = cell_text(row[KELG_WH_TITLE_COLUMN]) if len(row) > KELG_WH_TITLE_COLUMN else None
        con.execute(
            "INSERT INTO kelg_hymns VALUES (?,?,?,?,?)",
            (row[0], cell_text(row[1]), wh_title, cell_text(row[2]), supplement),
        )
        for code, col in KELG_REF_COLUMNS.items():
            raw = cell_text(row[col])
            if raw is None:
                continue
            verses = row[col + 2]
            if verses is not None and not isinstance(verses, (int, str)):
                verses = None  # a stray date-formatted cell
            con.execute(
                "INSERT INTO kelg_refs VALUES (?,?,?,?,?,NULL)",
                (row[0], code, raw, cell_text(row[col + 1]), cell_text(verses)),
            )


def load_wh(con):
    if not WH_CSV.exists():
        return
    with open(WH_CSV, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            con.execute(
                "INSERT INTO wh_hymns VALUES (?,?,?,?,?,?,?)",
                (int(r["number"]), r["english_first_line"] or None, r["german_first_line"] or None,
                 r["tune"] or None, r["tune_settings"] or None, r["first_stanza_line"] or None,
                 int(r["see_number"]) if r["see_number"] else None),
            )
            for code, setting, star in re.findall(r"\b(TLH|LW|LSB)\s*(\d+)\s*;?(\*?)", r["tune_settings"]):
                con.execute("INSERT INTO wh_tune_settings VALUES (?,?,?,?)",
                            (int(r["number"]), code, int(setting), int(bool(star))))


def load_tlhy(con):
    if not TLHY_CSV.exists():
        return
    with open(TLHY_CSV, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            con.execute(
                "INSERT INTO tlhy_index VALUES (?,?,?,?,?,?)",
                (int(r["number"]), r["first_line"], r["translator"] or None, int(r["alternate"] or 0),
                 int(r["index_page"]) if r["index_page"] else None, r["note"] or None),
            )


def read_csv(path):
    if not path.exists():
        return []
    with open(path, newline="", encoding="utf-8") as f:
        return [{k: ((v or "").strip() or None) for k, v in r.items()} for r in csv.DictReader(f)]


def load_tlh(con):
    for r in read_csv(TLH_CSV):
        con.execute("INSERT INTO tlh_hymns VALUES (?,?,?,?,?,?,?)",
                    (int(r["number"]), r["title"], r["first_line"], r["author"], r["translator"],
                     r["original_title"], r["tune"]))


def load_tlhy_originals(con):
    for r in read_csv(TLHY_ORIGINALS_CSV):
        con.execute("INSERT INTO tlhy_originals VALUES (?,?,?,?,?,?,?,?)",
                    (int(r["number"]), r["language"], r["original_title"], r["original_author"],
                     r["confidence"], r["same_as"], r["source"], r["note"]))


def load_tl(con):
    for r in read_csv(TL_CSV):
        con.execute("INSERT INTO tl_index VALUES (?,?,?,?,?)",
                    (parse_number(r["number"]), r["title"], int(r["alternate"] or 0),
                     int(r["index_page"]) if r["index_page"] else None, r["note"]))


def load_tl_originals(con):
    for r in read_csv(TL_ORIGINALS_CSV):
        con.execute("INSERT INTO tl_originals VALUES (?,?,?,?,?,?,?,?)",
                    (parse_number(r["number"]), r["language"], r["original_title"], r["original_author"],
                     r["confidence"], r["same_as"], r["source"], r["note"]))


def load_elgbs_lpc(con):
    for r in read_csv(ELGBS_LPC_CSV):
        con.execute("INSERT INTO elgbs_lpc_hymns VALUES (?,?,?,?,?,?,?,?,?,?)",
                    (int(r["number"]), r["german_first_line"], r["english_first_line"], r["author"],
                     r["translator"], r["tune"], int(r["elgbs_page"]), int(r["lpc_page"]),
                     r["row_title"], r["note"]))


# Common English words, to tell the English half of a first line Hymnary gives in two
# languages ('Stille Nacht, heilige Nacht! (Silent night, holy night)') from the other.
ENGLISH_WORDS = set("""the and of to in our we us you your thy thee thou ye is be are all with for his he
his who come let on from this that was have has when how what now my me god lord holy jesus christ
love night day praise glory sing hear shall will her him it its as by at so no not one may
o oh glorious earth heaven light king spirit son father savior people peace grace""".split())


def english_share(text):
    words = normalize_title(text).split()
    return sum(w in ENGLISH_WORDS for w in words) / len(words) if words else 0


def hymnary_parts(first_line, split=True):
    """(English first line, first line in another language or None) from a first line
    as Hymnary gives it: variant readings in brackets dropped ('From all that [who]
    dwell below the skies'), and a first line followed by its translation in
    parentheses ('Kyrie eleison (Lord have mercy)', 'Silent night, holy night!
    (Stille Nacht, heilige Nacht!)') split, the more English half first. A
    parenthesis that is English too stays part of the line; where the halves
    are as English, the one with accented letters ('Vengo a ti, Jesús amado') is
    the other language. A comma gets its
    space ('Come,ye thankful people'), and a line loses a comma or colon at its
    end ('Onward, Christian soldiers,'). Not `split`, a parenthesis is part of the
    line (LH's 'Upon the Lord's birthday bright (II)')."""
    text = " ".join(re.sub(r"\[[^\]]*\]", " ", first_line).split())
    text = re.sub(r"\s+([,;:!?.])", r"\1", text)
    text = re.sub(r",(?=[^\W\d_])", ", ", text)
    m = split and re.fullmatch(r"(.*?)\s*\(([^()]*)\)?\s*", text)
    if m and m.group(1):
        a, b = (re.sub(r"[,;:]+$", "", part.strip()) for part in m.groups())
        share_a, share_b = english_share(a), english_share(b)
        if share_a == share_b:  # a tie: the half with letters English lacks is the other language
            share_a -= bool(re.search(r"[À-ÖØ-öø-ɏ]", a))
            share_b -= bool(re.search(r"[À-ÖØ-öø-ɏ]", b))
        if share_b > share_a:
            return b, a
        if share_b <= 0 < share_a:
            return a, b
    return re.sub(r"[,;:]+$", "", text), None


def load_hymnary(con):
    """The Hymnary.org tables. Numbers outside the book's numbering are left out: the
    doxologies of a book's own doxology section (TBW07's D1-D28)."""
    unused = dict(HYMNARY_FIXES)
    for code, path in HYMNARY_CSVS.items():
        for r in read_csv(path):
            if not NUMBER_RE.match(r["number"]):
                continue
            first_line = r["first_line"]
            fix = unused.pop((code, r["number"]), None)
            if fix:
                if first_line != fix[0]:
                    raise SystemExit(f"HYMNARY_FIXES: {code} {r['number']} is now {first_line!r}, not {fix[0]!r}")
                first_line = fix[1]
            given = r.get("original")  # LH gives its Latin apart
            english, original = hymnary_parts(first_line, split=not given)
            original = given or original
            con.execute("INSERT INTO hymnary_index VALUES (?,?,?,?,?,?,?,?)",
                        (code, parse_number(r["number"]), first_line, english, original,
                         r.get("tune"), int(r.get("alternate") or 0), r.get("note")))
    if unused:
        raise SystemExit(f"HYMNARY_FIXES: no such hymn: {', '.join(f'{c} {n}' for c, n in unused)}")


def load_kelg_english(con):
    for r in read_csv(KELG_ENGLISH_CSV):
        con.execute("INSERT INTO kelg_english VALUES (?,?,?,?,?,?)",
                    (int(r["kelg_number"]), r["title"], r["translator"], r["source"],
                     r["source_url"], r["note"]))


def load_elhb(con):
    ws = openpyxl.load_workbook(ELHB_XLSX, read_only=True).worksheets[0]
    for row in ws.iter_rows(values_only=True):
        if row[0] is None and row[1] is None:
            continue
        con.execute("INSERT INTO elhb_hymns VALUES (?,?)",
                    (row[1], ELHB_INDEX_FIXES.get(row[1], cell_text(row[0]))))


def fix_kelg_elhb_refs(con):
    """Correct ELHB numbers in the KELG index whose title belongs to another ELHB hymn.

    Only single-number cells whose title matches exactly one other ELHB hymn
    are corrected. Check the ELHB itself before trusting a correction: the
    disagreements this used to correct (ELHB 230/286 and 431/461) were the
    ELHB index's mistakes, now fixed in ELHB_INDEX_FIXES.
    """
    elhb_titles = dict(con.execute("SELECT number, title FROM elhb_hymns"))
    for kelg_number, raw, title in con.execute(
        "SELECT kelg_number, raw, title FROM kelg_refs WHERE hymnal_code = 'ELHB'"
    ).fetchall():
        parsed = parse_cell(raw)
        if not title or len(parsed) != 1 or parsed[0][2]:
            continue
        n = parsed[0][0]
        if n in elhb_titles and titles_agree(title, elhb_titles[n]):
            continue
        guess = [m for m, t in elhb_titles.items() if titles_agree(title, t)]
        if len(guess) == 1:
            con.execute(
                "UPDATE kelg_refs SET fixed = ? WHERE kelg_number = ? AND hymnal_code = 'ELHB'",
                (str(guess[0]), kelg_number),
            )


def kelg_elhb_refs(con):
    """[(kelg_number, number cell, title, note)] using corrected ELHB numbers."""
    return [
        (k, fixed or raw, title,
         f"; KELG index gives ELHB {raw}, corrected by its title '{title}'" if fixed else "")
        for k, raw, fixed, title in con.execute(
            "SELECT kelg_number, raw, fixed, title FROM kelg_refs WHERE hymnal_code = 'ELHB'"
        ).fetchall()
    ]


# --- linking ---------------------------------------------------------------

def row_titles(con):
    """{hymn_id: [HYMN TITLE, alternative titles...]}: the titles a row bears."""
    out = defaultdict(list)
    for hymn_id, title in con.execute("SELECT id, title FROM hymns ORDER BY id"):
        out[hymn_id].append(title)
    for hymn_id, title in con.execute(
        "SELECT hymn_id, title FROM hymn_alternative_titles ORDER BY hymn_id, position"
    ):
        out[hymn_id].append(title)
    return out


def hymn_numbers_by_hymnal(con):
    """{hymn_id: {hymnal_code: {number, ...}}} for real (non-page) numbers; a
    number in a LETTERED_HYMNALS column keeps its letter ('157a')."""
    out = defaultdict(lambda: defaultdict(set))
    for hymn_id, code, number, suffix in con.execute(
        "SELECT hymn_id, hymnal_code, number, suffix FROM hymn_numbers WHERE is_page = 0"
    ):
        out[hymn_id][code].add(link_number(number, suffix) if code in LETTERED_HYMNALS else number)
    return out


def link_kelg(con, add_link, review):
    """Find each KELG hymn's master-index rows through its TLH/LW/ELH/LSB numbers.

    A row is a candidate when it shares a number with the KELG hymn in any of
    those hymnals. A candidate is linked 'ok' when it contradicts none of the
    KELG hymn's other numbers, and 'review' when it does (it carries a
    different number in a hymnal where the KELG index gives one).
    Returns {kelg_number: [hymn_id, ...]} for the links made.
    """
    numbers = hymn_numbers_by_hymnal(con)
    by_number = defaultdict(set)
    for hymn_id, per_hymnal in numbers.items():
        for code, nums in per_hymnal.items():
            for n in nums:
                by_number[code, n].add(hymn_id)

    refs = defaultdict(dict)
    for kelg_number, code, raw in con.execute("SELECT kelg_number, hymnal_code, raw FROM kelg_refs"):
        refs[kelg_number][code] = {n for n, _, page in parse_cell(raw) if not page}

    linked = defaultdict(list)
    for kelg_number, german_title in con.execute(
        "SELECT number, german_title FROM kelg_hymns ORDER BY number"
    ).fetchall():
        kref = {c: refs[kelg_number][c] for c in KELG_MATCH_HYMNALS if refs[kelg_number].get(c)}
        if not kref:
            continue  # no English counterpart in TLH/LW/ELH/LSB
        candidates = set()
        for code, nums in kref.items():
            for n in nums:
                candidates |= by_number[code, n]
        if not candidates:
            review.append(("KELG", kelg_number, german_title,
                           "no master-index row has its " + fmt_refs(kref)))
            continue
        for hymn_id in sorted(candidates):
            have = numbers[hymn_id]
            agree = sorted(c for c in kref if have.get(c, set()) & kref[c])
            adjacent = sorted(c for c in kref if have.get(c) and not have[c] & kref[c]
                              and not contradicts(have[c], kref[c]))
            disagree = sorted(c for c in kref if have.get(c) and contradicts(have[c], kref[c]))
            detail = "matched on " + ", ".join(agree)
            if adjacent:
                detail += "; another setting on " + ", ".join(
                    f"{c} (KELG index {fmt_nums(kref[c])}, master {fmt_nums(have[c])})" for c in adjacent
                )
            if disagree:
                detail += "; differs on " + ", ".join(
                    f"{c} (KELG index {fmt_nums(kref[c])}, master {fmt_nums(have[c])})" for c in disagree
                )
            status = "review" if disagree else "ok"
            add_link(hymn_id, "KELG", kelg_number, "hymnal numbers", status, detail)
            linked[kelg_number].append(hymn_id)
    return linked


def link_kelg_by_wh(con, kelg_links, add_link):
    """KELG hymns not found by hymnal number: look for master rows titled with
    the hymn's English first line in WH. Always 'review': WH's translation can
    share a first line with a different hymn."""
    by_norm = defaultdict(set)
    for hymn_id, titles in row_titles(con).items():
        for title in titles:
            by_norm[normalize_title(title)].add(hymn_id)
    have = set(con.execute("SELECT hymn_id, number FROM links WHERE hymnal_code = 'KELG'"))
    for number, english, first in con.execute(
        "SELECT number, english_first_line, first_stanza_line FROM wh_hymns"
    ).fetchall():
        if number in kelg_links:
            continue
        lines = [l for l in (english or "").split(" / ") + [first or ""] if l]
        for line in lines:
            norm = normalize_title(line)
            for key, ids in by_norm.items():
                if key == norm or (len(min(key, norm, key=len).split()) >= 4 and is_prefix(key, norm)):
                    for hymn_id in sorted(ids):
                        if (hymn_id, number) in have:
                            continue  # already linked; WH just agrees
                        add_link(hymn_id, "KELG", number, "WH first line", "review",
                                 f"WH {number} begins '{line}'")


# Spelling differences that shouldn't stop a TLHy first line matching a title.
ELISIONS = [("heav'n", "heaven"), ("o'er", "over"), ("ev'ry", "every"), ("pow'r", "power"),
            ("e'er", "ever"), ("ne'er", "never"), ("wher'er", "wherever"), ("whate'er", "whatever"),
            ("whene'er", "whenever"), ("hear'st", "hearest"), ("know'st", "knowest")]
# British spellings, as older and English hymnals print them ('Saviour').
BRITISH = {"saviour": "savior", "honour": "honor", "labour": "labor", "favour": "favor",
           "colour": "color", "splendour": "splendor", "harbour": "harbor", "valour": "valor",
           "ardour": "ardor", "vigour": "vigor", "clamour": "clamor", "endeavour": "endeavor",
           "neighbour": "neighbor", "odour": "odor", "rumour": "rumor", "succour": "succor",
           "behaviour": "behavior", "saviours": "saviors", "honours": "honors", "labours": "labors"}
# Archaic forms modernized in later hymnals ('Lord, Keep Us Steadfast in Your Word').
MODERN = {"thy": "your", "thee": "you", "thou": "you", "thine": "your", "ye": "you",
          "hast": "have", "art": "are", "doth": "does", "hath": "has", "didst": "did",
          "dost": "do", "wilt": "will", "shalt": "shall", "canst": "can"}
# Numbers that identify a hymn across a row's title variants, for spreading a
# TLHy number to the other rows of the same hymn. ELHB is left out: a few ELHB
# numbers in the master index are known to be wrong.
SAME_HYMN_HYMNALS = ("TLH", "LW", "LSB", "ELH", "KELG")


def match_key(title, modern=False):
    """normalize_title, plus: elisions spelled out, 'oh' -> 'o', British spellings
    American ('Saviour'), hyphenated words joined ('corner-stone'), and optionally
    archaic pronouns/verbs modernized."""
    text = title.lower().replace("\u2019", "'")
    for short, full in ELISIONS:
        text = text.replace(short, full)
    text = re.sub(r"(\w)'n\b", r"\1en", text)  # giv'n, ris'n
    text = re.sub(r"(\w)-(\w)", r"\1\2", text)
    words = ["o" if w == "oh" else BRITISH.get(w, w) for w in normalize_title(text).split()]
    if modern:
        words = [MODERN.get(w, w) for w in words]
    return " ".join(words)


def link_by_title(con, add_link, code, entries, one_row=False, rejected=frozenset()):
    """Match a hymnal's titles (TLHy's first lines, TL's index titles), given as
    [(number, title)], to the titles master-index rows are known by.

    A title equal to one of a row's titles (its HYMN TITLE, or its WH, ELHB,
    TLHy, TLH, LW, ELH, LSB title, or any English title the KELG index gives
    its KELG hymn) links 'ok', also after modernizing archaic words or
    ignoring word breaks ('Christmas tide', 'Christmastide'). The number then
    also goes on the other rows of the same hymn: rows sharing a KELG or TLH
    number with a matched row, or an LW, LSB or ELH number and a similar title
    ('ok'; other relatives 'review'). A title that only nearly matches, or is
    the start of a row's title (or the reverse), links 'review', unless it is
    the start of more than PREFIX_LIMIT rows' titles: too common a start to tell
    anything ('Lord Jesus Christ'). What is left is matched by its original next
    (link_by_original).

    With `one_row` (the Hymnary.org hymnals), a title equal to the titles of
    several rows that aren't one hymn (sharing a KELG, TLH, LW, LSB or ELH
    number) links 'review', since a first line two hymns share can't tell them
    apart ('Come, Holy Spirit'); a row whose link `rejected` holds (the (row,
    hymnal, number) links link_decisions.csv calls 'no match') doesn't count.
    """
    titles = list(con.execute("SELECT hymn_id, title FROM hymn_titles WHERE source != ?", (code,)))
    kelg_titles = defaultdict(list)
    for k, raw, t in con.execute("SELECT kelg_number, raw, title FROM kelg_refs WHERE title IS NOT NULL"):
        kelg_titles[k] += [title for _, title in ref_titles(raw, t)]
    for hymn_id, k in con.execute("SELECT hymn_id, number FROM hymn_numbers WHERE hymnal_code = 'KELG'"):
        titles += [(hymn_id, t) for t in kelg_titles[k]]
    exact, modern, unspaced = defaultdict(set), defaultdict(set), defaultdict(set)
    for hymn_id, t in titles:
        exact[match_key(t)].add(hymn_id)
        modern[match_key(t, modern=True)].add(hymn_id)
        unspaced[match_key(t, modern=True).replace(" ", "")].add(hymn_id)
    hymn_titles = dict(con.execute("SELECT id, title FROM hymns"))
    titles_of = row_titles(con)

    by_number = defaultdict(set)  # (code, number) -> rows
    row_numbers = defaultdict(set)
    for hymn_id, other, n in con.execute(
        "SELECT hymn_id, hymnal_code, number FROM hymn_numbers WHERE is_page = 0"
    ):
        if other in SAME_HYMN_HYMNALS:
            by_number[other, n].add(hymn_id)
            row_numbers[hymn_id].add((other, n))

    def one_hymn(rows):
        """Whether `rows` are joined by the numbers they share."""
        rows = set(rows)
        start = next(iter(rows))
        seen, todo = {start}, [start]
        while todo:
            for key in row_numbers[todo.pop()]:
                for h in (by_number[key] & rows) - seen:
                    seen.add(h)
                    todo.append(h)
        return seen == rows

    placed = defaultdict(set)  # number -> the rows the master index has it on
    for hymn_id, n, suffix in con.execute(
        "SELECT hymn_id, number, suffix FROM hymn_numbers WHERE hymnal_code = ? AND is_page = 0", (code,)
    ):
        placed[link_number(n, suffix) if code in LETTERED_HYMNALS else n].add(hymn_id)

    unmatched = defaultdict(list)
    direct_rows = defaultdict(set)  # row -> numbers matched by its own titles
    spread = defaultdict(lambda: defaultdict(set))  # row -> number -> {(code, n, matched row)}
    source_titles = {}
    for number, title in entries:
        source_titles.setdefault(number, title)
        found = {h: "title" for h in exact.get(match_key(title), ())}
        for h in modern.get(match_key(title, modern=True), ()):
            found.setdefault(h, "modernized title")
        if not found:
            for h in unspaced.get(match_key(title, modern=True).replace(" ", ""), ()):
                found.setdefault(h, "title, word breaks aside")
        if number in placed:  # a title only confirms the rows it is on (see build's add_link)
            for hymn_id, method in sorted(found.items()):
                if hymn_id in placed[number]:
                    add_link(hymn_id, code, number, method, "ok", f"'{title}'")
            continue
        if not found:
            unmatched[number].append(title)
            continue
        status, detail = "ok", f"'{title}'"
        live = [h for h in found if (h, code, number) not in rejected]
        if one_row and len(live) > 1 and not one_hymn(live):
            status, detail = "review", f"'{title}', which {len(live)} rows that aren't one hymn bear"
        for hymn_id, method in sorted(found.items()):
            add_link(hymn_id, code, number, method, status, detail)
            direct_rows[hymn_id].add(number)
            for other, n in row_numbers[hymn_id]:
                for sibling in by_number[other, n] - found.keys():
                    spread[sibling][number].add((other, n, hymn_id))

    # Spread to the other rows of the same hymn. A shared KELG or TLH number is
    # taken as the same hymn; a shared LW, LSB or ELH number only when the titles
    # are alike too, since those hymnals sometimes print several hymns under one
    # number. Rows that matched a hymn of this hymnal by their own title keep it,
    # and rows whose relatives matched several of its hymns are left for review.
    for hymn_id, found in sorted(spread.items()):
        if hymn_id in direct_rows:
            continue
        for number, shared in sorted(found.items(), key=lambda kv: number_order(kv[0])):
            other, n, matched = min(shared, key=lambda s: (s[0] not in ("KELG", "TLH"), s))
            why = f"'{source_titles[number]}' is '{hymn_titles[matched]}', which shares {other} {n}"
            strong = any(c in ("KELG", "TLH") for c, _, _ in shared)
            alike = max(difflib.SequenceMatcher(
                None, match_key(source_titles[number], modern=True), match_key(t, modern=True)
            ).ratio() for t in titles_of[hymn_id]) >= 0.8
            status = "ok" if strong or alike else "review"
            if status == "review":
                why += "; titles differ and only that number ties them"
            if len(found) > 1:
                status = "review"
                why += f"; rows sharing its numbers matched {code} " + fmt_nums(found)
            add_link(hymn_id, code, number, "same hymn", status, why)

    linked = {n for (n,) in con.execute("SELECT DISTINCT number FROM links WHERE hymnal_code = ?", (code,))}
    mkeys = list(modern)
    for number, unplaced in sorted(unmatched.items(), key=lambda kv: number_order(kv[0])):
        if number in linked:
            continue  # another index entry for this hymn matched
        for title in unplaced:
            mkey = match_key(title, modern=True)
            close = difflib.get_close_matches(mkey, mkeys, n=3, cutoff=FUZZY_THRESHOLD)
            prefixed = [k for k in mkeys if k not in close
                        and len(min(k, mkey, key=len).split()) >= 3 and is_prefix(k, mkey)]
            longer = [k for k in prefixed if len(k) > len(mkey)]
            if len(longer) > PREFIX_LIMIT:
                prefixed = [k for k in prefixed if k not in longer]
            for k in close + prefixed:
                how = (f"{difflib.SequenceMatcher(None, mkey, k).ratio():.0%} similar" if k in close
                       else "one is the start of the other")
                for hymn_id in sorted(modern[k]):
                    add_link(hymn_id, code, number, "similar title" if k in close else "title prefix",
                             "review", f"'{title}' vs '{hymn_titles[hymn_id]}' ({how})")


def link_tlhy(con, add_link):
    link_by_title(con, add_link, "TLHy", con.execute(
        "SELECT number, first_line FROM tlhy_index ORDER BY number, alternate").fetchall())


def natural_title(title):
    """An index title back in reading order: 'Great "O" Antiphons, The' ->
    'The Great "O" Antiphons'."""
    m = re.match(r"(.+), (The|A|An)$", title)
    return f"{m.group(2)} {m.group(1)}" if m else title


def tl_titles(con):
    """{TL number: [title, ...]}, main index entry first, in reading order."""
    out = defaultdict(list)
    for number, title in con.execute("SELECT number, title FROM tl_index ORDER BY alternate"):
        out[number].append(natural_title(title))
    return out


def link_tl(con, add_link):
    """Match TL's index titles; one with a parenthesis ('I Believe in One God
    (Nicene Creed)') also without it."""
    entries = []
    for number, titles in sorted(tl_titles(con).items(), key=lambda kv: number_order(kv[0])):
        for title in titles:
            entries.append((number, title))
            bare = re.sub(r"\s*\([^)]*\)", "", title)
            if bare != title:
                entries.append((number, bare))
    link_by_title(con, add_link, "TL", entries)


def link_lpc(con, add_link):
    """LPC's hymns, by LPC's English first line, then by the German original
    (ELGBS, which numbers its house hymnal as LPC does). ELGBS's house hymnal
    was made of hymns not in KELG, and some open like a KELG hymn (35, 'Lobt
    Gott, ihr Christen allzugleich'), so a match through KELG wants a look
    before it is trusted (link_decisions.csv)."""
    link_by_title(con, add_link, "LPC", con.execute(
        "SELECT number, english_first_line FROM elgbs_lpc_hymns ORDER BY number").fetchall())
    link_by_original(con, add_link, "LPC", "lpc_originals")


def hymnary_titles(con, code):
    """{number: [English first line, ...]} of a Hymnary.org hymnal: Hymnary's first,
    then the others the book's own index gives (WS69)."""
    out = defaultdict(list)
    for number, english in con.execute(
        "SELECT number, english FROM hymnary_index WHERE hymnal = ? ORDER BY alternate, rowid", (code,)
    ):
        out[number].append(english)
    return out


def hymnary_codes(con):
    """The Hymnary.org hymnals the master index has a column for, in HYMNARY_CSVS order."""
    present = {c for (c,) in con.execute("SELECT code FROM hymnals")}
    return [code for code in HYMNARY_CSVS if code in present]


def link_hymnary(con, add_link, decisions):
    """The Hymnary.org hymnals' hymns, one hymnal after the other, by their English
    first lines and, where Hymnary gives one, the first line in the other language.
    A first line several hymns share is left for review (link_by_title)."""
    rejected = {(h, code, n) for h, code, n, decision, _ in decisions if decision == "no match"}
    for code in hymnary_codes(con):
        entries = []
        for number, english, original in con.execute(
            "SELECT number, english, original FROM hymnary_index WHERE hymnal = ? ORDER BY alternate, rowid",
            (code,)
        ).fetchall():
            entries.append((number, english))
            if original:
                entries.append((number, original))
        entries.sort(key=lambda e: number_order(e[0]))
        link_by_title(con, add_link, code, entries, one_row=True, rejected=rejected)


def translator_names(credit):
    """Surnames in a translator credit, spelled loosely: 'Miss C. Winkworth, 1863'
    -> {'winkworth'}; 'Carl Doeving, 1909' and 'C. Doving' -> {'doving'}."""
    names = set()
    for part in re.split(r"[,;&]|\band\b", credit or ""):
        words = [w for w in re.findall(r"[^\W\d_]{3,}", part)
                 if w.lower() not in ("miss", "mrs", "alt", "composite", "unknown", "cento")]
        if words:
            name = normalize_title(words[-1].replace("ø", "o")).replace(" ", "")
            names.add(re.sub(r"([aou])e", r"\1", name))
    return names


def confirm_tlhy_by_translator(con):
    """A TLHy first line that only nearly matches a row's title is confirmed
    when the row's TLH hymn has the same translator (TLH 1941 reprints many
    of TLHy's translations, lightly altered)."""
    tlh = {n: translator_names(t) for n, t in con.execute(
        "SELECT number, translator FROM tlh_hymns WHERE translator IS NOT NULL")}
    tlhy = {n: translator_names(t) for n, t in con.execute(
        "SELECT number, translator FROM tlhy_index WHERE alternate = 0 AND translator IS NOT NULL")}
    row_tlh = defaultdict(set)
    for hymn_id, n in con.execute(
        "SELECT hymn_id, number FROM hymn_numbers WHERE hymnal_code = 'TLH' AND is_page = 0"
    ):
        row_tlh[hymn_id].add(n)
    for hymn_id, number in con.execute(
        "SELECT hymn_id, number FROM links WHERE hymnal_code = 'TLHy' AND status = 'review' "
        "AND method IN ('similar title', 'title prefix')"
    ).fetchall():
        same = [n for n in row_tlh[hymn_id] if tlh.get(n, set()) & tlhy.get(number, set())]
        if same:
            con.execute(
                "UPDATE links SET status = 'ok', detail = detail || ? "
                "WHERE hymn_id = ? AND hymnal_code = 'TLHy' AND number = ?",
                (f"; confirmed: TLH {fmt_nums(same)} has the same translator", hymn_id, number),
            )


ORIGINAL_FOLDS = (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss"), ("ø", "oe"), ("æ", "ae"),
                  ("å", "aa"), ("œ", "oe"), ("'", ""), ("’", ""))


def original_key(title):
    """normalize_title for an original first line, with spellings folded:
    umlauts and their ae/oe/ue form (tlh_1941.csv is ASCII), ß/ss, elisions
    ('Sünd'' = 'Sünd'), and old th/t, y/i, c/k, j/i spellings."""
    text = title.lower()
    for a, b in ORIGINAL_FOLDS:
        text = text.replace(a, b)
    text = normalize_title(text)
    for a, b in (("th", "t"), ("y", "i"), ("c", "k"), ("j", "i")):
        text = text.replace(a, b)
    return text


def original_matches(key, keys, min_words):
    """The numbers in {number: original_key()} whose first line is `key`: equal
    (word breaks aside); failing that, one starts the other and the shorter has
    at least min_words words; failing that, at least 95% alike (a misprint)."""
    flat = key.replace(" ", "")
    for agree in (lambda k: k.replace(" ", "") == flat,
                  lambda k: len(min(k, key, key=len).split()) >= min_words and is_prefix(k, key),
                  lambda k: difflib.SequenceMatcher(None, k, key).ratio() >= 0.95):
        hits = [n for n, k in keys.items() if agree(k)]
        if hits:
            return hits
    return []


def link_by_original(con, add_link, code, table):
    """Line up a hymnal's translations by their originals (`table`: tlhy_originals
    for TLHy, tl_originals for TL).

    Where the research found the very same translation under another
    hymnal's number ('same_as', e.g. 'ELHB 456'), it goes on that row.
    Otherwise one whose original is a KELG hymn goes on that hymn's rows; failing
    that, one whose original is a TLH hymn's (the TLH 1941 notes; they shorten
    some first lines, hence the shorter match) goes on that TLH hymn's rows;
    failing that, it goes on the rows whose Original Title it is.
    'ok' when the original is confirmed or probable and every hymn with that
    first line is on the same rows (different hymns can share a first line);
    otherwise 'review'. A near-match link to the same row is confirmed. Hymns
    whose own title already found their rows are left alone: other
    translations of the same original usually have rows of their own.
    """
    hymn_titles = dict(con.execute("SELECT id, title FROM hymns"))
    kelg_rows, tlh_rows = defaultdict(set), defaultdict(set)
    for hymn_id, other, n in con.execute(
        "SELECT hymn_id, hymnal_code, number FROM hymn_numbers WHERE hymnal_code IN ('KELG', 'TLH') AND is_page = 0"
    ):
        (kelg_rows if other == "KELG" else tlh_rows)[n].add(hymn_id)
    for hymn_id, n in con.execute("SELECT hymn_id, number FROM links WHERE hymnal_code = 'KELG' AND status = 'ok'"):
        kelg_rows[n].add(hymn_id)
    originals = dict(con.execute("SELECT id, original_title FROM hymns WHERE original_title IS NOT NULL"))
    sources = (  # (where, {key: original}, {key: rows}, shortest prefix match, what the key is)
        ("KELG", dict(con.execute("SELECT number, german_title FROM kelg_hymns")), kelg_rows, 3,
         lambda hits: f"KELG {fmt_nums(hits)}"),
        ("TLH", dict(con.execute("SELECT number, original_title FROM tlh_hymns WHERE original_title IS NOT NULL")),
         tlh_rows, 2, lambda hits: f"TLH {fmt_nums(hits)}'s original in the TLH 1941 notes"),
        ("Original Title", originals, {h: {h} for h in originals}, 3,
         lambda hits: "the Original Title of " + ", ".join(f"'{hymn_titles[h]}'" for h in sorted(hits))),
    )
    keys = {where: {n: original_key(t) for n, t in titles.items()} for where, titles, _, _, _ in sources}
    by_number = defaultdict(set)
    for hymn_id, other, n in con.execute("SELECT hymn_id, hymnal_code, number FROM hymn_numbers WHERE is_page = 0"):
        by_number[other, n].add(hymn_id)
    prior = {(h, n): (status, method) for h, n, status, method in con.execute(
        "SELECT hymn_id, number, status, method FROM links WHERE hymnal_code = ?", (code,))}

    def link(number, rows, method, status, detail):
        if status == "ok":
            for (hymn_id, n), was in prior.items():
                if n == number and hymn_id not in rows and was[0] == "review" and was[1] != "same hymn":
                    con.execute(
                        "UPDATE links SET detail = detail || ? "
                        "WHERE hymn_id = ? AND hymnal_code = ? AND number = ?",
                        (f"; probably wrong: its original puts it elsewhere ({detail})", hymn_id, code, number),
                    )
        for hymn_id in sorted(rows):
            was = prior.get((hymn_id, number))
            if was is None:
                add_link(hymn_id, code, number, method, status, detail)
            elif was in (("review", "similar title"), ("review", "title prefix")) and status == "ok":
                con.execute(
                    "UPDATE links SET status = 'ok', detail = detail || ? "
                    "WHERE hymn_id = ? AND hymnal_code = ? AND number = ?",
                    (f"; confirmed: {detail}", hymn_id, code, number),
                )

    matched = {n for (n,) in con.execute(
        "SELECT number FROM links WHERE hymnal_code = ? AND status = 'ok'", (code,))}
    for number, title, author, confidence, same_as in con.execute(
        f"SELECT number, original_title, author, confidence, same_as FROM {table} "
        "WHERE original_title IS NOT NULL OR same_as IS NOT NULL"
    ).fetchall():
        if number in matched:
            continue  # its own title already found its rows
        status = "ok" if confidence in TRUSTED else "review"
        translates = f"translates '{title}'" + (f" ({author})" if author else "")
        if same_as:
            other, n = same_as.split()
            rows = by_number[other, parse_number(n)]
            if not rows:
                raise SystemExit(f"{table}: {code} {number} same_as {same_as}, but no row has {same_as}")
            link(number, rows, "same translation", status, f"{translates}; the same translation as {same_as}")
            continue
        key = original_key(title)
        for where, titles, rows_of, min_words, describe in sources:
            hits = sorted(n for n in original_matches(key, keys[where], min_words) if rows_of.get(n))
            if not hits:
                continue
            rows = set().union(*(rows_of[n] for n in hits))
            detail = f"{translates}, {describe(hits)}"
            if confidence not in TRUSTED:
                detail += "; original not confirmed"
            elif where != "Original Title" and not all(rows_of[n] == rows for n in hits):
                status = "review"
                detail += "; different hymns share that first line"
            link(number, rows, f"original in {where}", status, detail)
            break


def report_unmatched_tlhy(con, review):
    linked = {n for (n,) in con.execute(
        "SELECT DISTINCT number FROM links WHERE hymnal_code = 'TLHy' AND status != 'rejected'")}
    originals = dict(con.execute(
        "SELECT number, language FROM tlhy_originals WHERE language IS NOT NULL"))
    for number, first_line, translator in con.execute(
        "SELECT number, first_line, translator FROM tlhy_index WHERE alternate = 0 ORDER BY number"
    ).fetchall():
        if number in linked:
            continue
        language = originals.get(number)
        if language and language != "English":
            why = f"a hymn new to the master index, translated from {language}"
        elif translator and language != "English":
            why = "a hymn new to the master index; a translation whose original wasn't traced"
        else:
            why = "a hymn new to the master index (English original)"
        review.append(("TLHy", number, first_line, "no matching English title or original: " + why))


def report_unmatched_tl(con, review):
    linked = {n for (n,) in con.execute(
        "SELECT DISTINCT number FROM links WHERE hymnal_code = 'TL' AND status != 'rejected'")}
    originals = dict(con.execute("SELECT number, language FROM tl_originals WHERE language IS NOT NULL"))
    for number, titles in sorted(tl_titles(con).items(), key=lambda kv: number_order(kv[0])):
        if number in linked:
            continue
        language = originals.get(number)
        if language == "English":
            why = "a hymn new to the master index (English original)"
        elif language:
            why = f"a hymn new to the master index, translated from {language}"
        else:
            why = "a hymn new to the master index; original not traced"
        review.append(("TL", number, titles[0], "no matching English title or original: " + why))


def report_unmatched_lpc(con, review):
    linked = {n for (n,) in con.execute(
        "SELECT DISTINCT number FROM links WHERE hymnal_code = 'LPC' AND status != 'rejected'")}
    for number, english, german in con.execute(
        "SELECT number, english_first_line, german_first_line FROM elgbs_lpc_hymns ORDER BY number"
    ).fetchall():
        if number not in linked:
            review.append(("LPC", number, english, "no matching English title or original: a hymn new to "
                           f"the master index, translated from German ('{german}')"))


def report_unmatched_hymnary(con, review):
    for code in hymnary_codes(con):
        linked = {n for (n,) in con.execute(
            "SELECT DISTINCT number FROM links WHERE hymnal_code = ? AND status != 'rejected'", (code,))}
        for number, titles in sorted(hymnary_titles(con, code).items(), key=lambda kv: number_order(kv[0])):
            if number not in linked:
                review.append((code, number, titles[0],
                               "no matching English title: a hymn new to the master index"))


def link_elhb(con, kelg_links, placed, add_link, review):
    """Link ELHB hymns to master-index rows.

    1. Through KELG: a KELG hymn's ELHB number goes on the rows that KELG hymn
       was linked to.
    2. By title: the ELHB title matches a master-index title apart from
       capitalization, punctuation, elisions and Thee/You ('ok'), or closely
       enough by fuzzy match ('review').
    """
    elhb_titles = dict(con.execute("SELECT number, title FROM elhb_hymns"))
    titles_of = row_titles(con)
    hymn_titles = {hymn_id: titles[0] for hymn_id, titles in titles_of.items()}
    keys_of = {hymn_id: {title_key(t) for t in titles} for hymn_id, titles in titles_of.items()}
    by_norm, by_flat = defaultdict(list), defaultdict(list)
    for hymn_id, keys in keys_of.items():
        for key in sorted(keys):
            by_norm[key].append(hymn_id)
            if hymn_id not in by_flat[key.replace(" ", "")]:
                by_flat[key.replace(" ", "")].append(hymn_id)
    norm_keys = list(by_norm)

    found = defaultdict(set)  # elhb number -> hymn ids
    for (code, n), ids in placed.items():
        if code == "ELHB":
            found[n] |= ids

    for kelg_number, raw, ref_title, note in kelg_elhb_refs(con):
        for n, _, page in parse_cell(raw):
            if page:
                continue
            index_title = elhb_titles.get(n)
            if index_title is None:
                note += f"; ELHB {n} is not in the ELHB index"
            elif ref_title and not titles_agree(ref_title, index_title):
                note = f"; KELG index calls ELHB {n} '{ref_title}' but ELHB index says '{index_title}'"
                guess = [m for m, t in elhb_titles.items() if titles_agree(ref_title, t)]
                if guess:
                    note += " (ELHB index has that title at " + fmt_nums(guess) + ")"
            for hymn_id in kelg_links.get(kelg_number, []):
                add_link(hymn_id, "ELHB", n, "via KELG", "review" if note else "ok",
                         f"KELG {kelg_number}" + note)
                found[n].add(hymn_id)

    via_kelg = {n: set(ids) for n, ids in found.items()}
    for n, title in sorted(elhb_titles.items()):
        norm = title_key(title)
        same = by_norm.get(norm) or by_flat.get(norm.replace(" ", ""))  # word breaks aside
        if same:
            kelg_rows = via_kelg.get(n, set())
            kelg_norms = set().union(*(keys_of[h] for h in kelg_rows)) if kelg_rows else set()
            for hymn_id in same:
                status, detail = "ok", "same title"
                if kelg_rows and hymn_id not in kelg_rows and not keys_of[hymn_id] & kelg_norms:
                    status = "review"
                    detail += "; KELG route linked ELHB {} to {}".format(
                        n, ", ".join(f"'{hymn_titles[h]}'" for h in sorted(kelg_rows)))
                add_link(hymn_id, "ELHB", n, "title", status, detail)
                found[n].add(hymn_id)
            continue
        if found[n]:
            continue
        close = difflib.get_close_matches(norm, norm_keys, n=3, cutoff=FUZZY_THRESHOLD)
        for key in close:
            ratio = difflib.SequenceMatcher(None, norm, key).ratio()
            for hymn_id in by_norm[key]:
                add_link(hymn_id, "ELHB", n, "similar title", "review",
                         f"'{title}' vs '{hymn_titles[hymn_id]}' ({ratio:.0%} similar)")
                found[n].add(hymn_id)
        if close:
            continue
        # One title is the start of the other, e.g. 'Abide with me; fast falls
        # the eventide' vs 'Abide with Me', or a truncated ELHB index entry.
        prefixed = [k for k in norm_keys if len(min(k, norm, key=len).split()) >= 3 and is_prefix(k, norm)]
        for key in prefixed:
            for hymn_id in by_norm[key]:
                add_link(hymn_id, "ELHB", n, "title prefix", "review",
                         f"'{title}' vs '{hymn_titles[hymn_id]}'")
                found[n].add(hymn_id)
        if not prefixed:
            review.append(("ELHB", n, title, "no master-index row found"))
    return found


def link_kelg_via_elhb(con, elhb_found, add_link):
    """KELG hymns whose only English counterpart is in ELHB get the master rows
    that ELHB hymn was linked to by title."""
    statuses = {(h, n): st for h, n, st in con.execute(
        "SELECT hymn_id, number, status FROM links WHERE hymnal_code = 'ELHB'")}
    linked = {n for (n,) in con.execute("SELECT number FROM links WHERE hymnal_code = 'KELG'")}
    for kelg_number, raw, _, note in kelg_elhb_refs(con):
        if kelg_number in linked:
            continue
        for n, _, page in parse_cell(raw):
            if page:
                continue
            for hymn_id in sorted(elhb_found.get(n, ())):
                add_link(hymn_id, "KELG", kelg_number, "via ELHB",
                         "review" if note else statuses[hymn_id, n],
                         f"KELG index gives ELHB {n}, which was linked by title" + note)


def is_prefix(a, b):
    short, long_ = sorted((a, b), key=len)
    return long_ == short or long_.startswith(short + " ")


def title_key(title):
    """Titles that differ only in capitalization, punctuation, spelling of
    elisions (ev'ry), archaic pronouns and verbs (Thee/You) or 'O'/'Oh' are
    the same title."""
    return match_key(title, modern=True)


def titles_agree(a, b):
    a, b = title_key(a), title_key(b)
    return (a.replace(" ", "") == b.replace(" ", "") or is_prefix(a, b)
            or difflib.SequenceMatcher(None, a, b).ratio() >= FUZZY_THRESHOLD)


def contradicts(have, want):
    """Two sets of numbers for one hymnal disagree. Numbers next to each other
    (LW 297/298 'A Mighty Fortress') are usually two settings of one hymn and
    don't count."""
    return not any(abs(a - b) <= 1 for a in have for b in want)


def fmt_nums(nums):
    return ", ".join(str(n) for n in sorted(nums, key=number_order))


def fmt_refs(kref):
    return ", ".join(f"{c} {fmt_nums(nums)}" for c, nums in kref.items())


def ref_titles(raw, title):
    """A KELG-index number cell and its title as [(numbers, title)], a cell
    naming two hymns split in two: '569 + 535' with 'O Lord, I Sing with Lips
    and Heart + Rejoice, My Heart, Be Glad and Sing', or '333' and '336' on two
    lines."""
    nums = [n for n, _, page in parse_cell(raw) if not page]
    parts = [t.strip() for t in re.split(r"\n| \+ ", title) if t.strip()]
    if len(parts) > 1 and len(parts) == len(nums):
        return [({n}, t) for n, t in zip(nums, parts)]
    return [(set(nums), title.strip())]


def build_hymn_titles(con):
    """Collect each master-index row's titles: its HYMN TITLE and Alternative
    English Titles, and from the numbers in its own cells: KELG -> German title
    and WH title; ELHB -> ELHB index title; TLHy -> TLHy first line; TL -> TL
    index title; LPC -> LPC's English and ELGBS's German first line; and, for a
    row with a KELG number, the KELG index's TLH/LW/ELH/LSB title wherever the
    row has that same number."""
    numbers = hymn_numbers_by_hymnal(con)
    con.execute("INSERT INTO hymn_titles SELECT id, 'master', title FROM hymns")
    con.execute("INSERT OR IGNORE INTO hymn_titles SELECT hymn_id, 'alternate', title FROM hymn_alternative_titles")
    kelg = {n: (g, w) for n, g, w in con.execute("SELECT number, german_title, wh_title FROM kelg_hymns")}
    elhb = {n: ELHB_TITLE_FIXES.get(n, t) for n, t in con.execute("SELECT number, title FROM elhb_hymns")}
    tlh = dict(con.execute("SELECT number, title FROM tlh_hymns"))
    refs = defaultdict(dict)
    for k, code, raw, title in con.execute("SELECT kelg_number, hymnal_code, raw, title FROM kelg_refs"):
        if title and code in KELG_MATCH_HYMNALS:
            refs[k][code] = ref_titles(raw, title)
    tlhy = defaultdict(list)
    for n, first_line in con.execute("SELECT number, first_line FROM tlhy_index ORDER BY alternate"):
        tlhy[n].append(first_line)
    tl = tl_titles(con)
    house = {n: (e, g) for n, e, g in con.execute(
        "SELECT number, english_first_line, german_first_line FROM elgbs_lpc_hymns")}
    hymnary = {code: hymnary_titles(con, code) for code in HYMNARY_CSVS}
    rows = []
    for hymn_id, have in numbers.items():
        for n in have.get("ELHB", ()):
            if n in elhb:
                rows.append((hymn_id, "ELHB", elhb[n]))
        for n in have.get("TLHy", ()):
            rows += [(hymn_id, "TLHy", t) for t in tlhy[n]]
        for n in sorted(have.get("TL", ()), key=number_order):
            rows += [(hymn_id, "TL", t) for t in tl[n]]
        rows += [(hymn_id, "LPC", house[n][0]) for n in sorted(have.get("LPC", ())) if n in house]
        rows += [(hymn_id, "ELGBS", house[n][1]) for n in sorted(have.get("LPC", ())) if n in house]
        for code, titles in hymnary.items():
            for n in sorted(have.get(code, ()), key=number_order):
                rows += [(hymn_id, code, t) for t in titles.get(n, ())]
        for k in have.get("KELG", ()):
            german, wh = kelg.get(k, (None, None))
            rows += [(hymn_id, "KELG", german)] if german else []
            # WH sometimes prints two English first lines: 'This body in the grave we lay / Ye bear me ...'
            rows += [(hymn_id, "WH", w.strip()) for w in (wh or "").split(" / ") if w.strip()]
            for code, parts in refs[k].items():
                for nums, title in parts:
                    if have.get(code, set()) & nums:
                        rows.append((hymn_id, code, title))
        rows += [(hymn_id, "TLH", tlh[n]) for n in sorted(have.get("TLH", ())) if n in tlh]
    # One title per spelling: the KELG index's 'Open now Thy gates of beauty'
    # and TLH 1941's 'Open Now Thy Gates of Beauty' are the same title.
    seen = set()
    for hymn_id, source, title in rows:
        key = (hymn_id, source, match_key(title))
        if key not in seen:
            seen.add(key)
            con.execute("INSERT OR IGNORE INTO hymn_titles VALUES (?,?,?)", (hymn_id, source, title))
    sources = ("KELG", "WH", "ELHB", "TLHy", "TLH", "LW", "ELH", "LSB", "TL", "LPC", "ELGBS", *HYMNARY_CSVS)
    pivot = ",\n    ".join(
        f"(SELECT group_concat(title, ' / ') FROM hymn_titles t WHERE t.hymn_id = h.id "
        f"AND t.source = '{src}') AS {src}_title"
        for src in sources
    )
    con.execute(f"""CREATE VIEW hymn_titles_by_hymnal AS
SELECT h.id, h.title AS master_title,
    {pivot}
FROM hymns h ORDER BY h.id""")
    codes = [c for (c,) in con.execute("SELECT code FROM hymnals ORDER BY sort_order")]
    con.execute(f"""CREATE VIEW hymn_catalog AS
SELECT i.id,
    {", ".join(f'i."{c}"' for c in codes)},
    i.german_title,
    t.master_title, i.alternative_titles, {", ".join(f"t.{src}_title" for src in sources if src != "KELG")},
    i.original_title, i.original_language, i.other_english_title, i.other_english_source,
    i.scripture
FROM hymn_index i JOIN hymn_titles_by_hymnal t ON t.id = i.id
ORDER BY i.id""")


def link_from_master(con, add_link):
    """KELG/ELHB/TLHy/TL/LPC/WS numbers already in the master index are links in their own right.

    An ELHB number is flagged when the ELHB index gives that number a different
    title; a KELG number when the KELG index gives that hymn a TLH/LW/ELH/LSB
    number the row contradicts. Returns {(code, number): {hymn_id, ...}}.
    """
    numbers = hymn_numbers_by_hymnal(con)
    elhb_titles = dict(con.execute("SELECT number, title FROM elhb_hymns"))
    titles_of = row_titles(con)
    refs = defaultdict(dict)
    for kelg_number, code, raw in con.execute("SELECT kelg_number, hymnal_code, raw FROM kelg_refs"):
        refs[kelg_number][code] = {n for n, _, page in parse_cell(raw) if not page}
    placed = defaultdict(set)
    for hymn_id, title, code, number, suffix in con.execute(
        """SELECT h.id, h.title, n.hymnal_code, n.number, n.suffix FROM hymn_numbers n
           JOIN hymns h ON h.id = n.hymn_id
           WHERE n.hymnal_code IN ({}) AND n.is_page = 0""".format(", ".join(f"'{c}'" for c in LINKED_HYMNALS))
    ).fetchall():
        if code in LETTERED_HYMNALS:
            number = link_number(number, suffix)
        status, detail = "ok", "already in master index"
        if code == "ELHB":
            index_title = elhb_titles.get(number)
            if index_title is None:
                status, detail = "review", f"ELHB {number} is not in the ELHB index"
            elif not any(titles_agree(t, index_title) for t in titles_of[hymn_id]):
                status = "review"
                detail += f"; ELHB index calls ELHB {number} '{index_title}'"
        elif code == "KELG":
            have = numbers[hymn_id]
            bad = [c for c in KELG_MATCH_HYMNALS
                   if refs[number].get(c) and have.get(c) and contradicts(have[c], refs[number][c])]
            if bad:
                status = "review"
                detail += "; differs from KELG index on " + ", ".join(
                    f"{c} (KELG index {fmt_nums(refs[number][c])}, master {fmt_nums(have[c])})" for c in bad)
        add_link(hymn_id, code, number, "in master", status, detail)
        placed[code, number].add(hymn_id)
    return placed


def read_same_hymn_decisions():
    """same_hymn_decisions.csv as [(title, master_title, decision, identified_by)].
    'same hymn': the row titled `title` is merged into the row titled
    `master_title` (write_back.py), which keeps its title; 'different hymns':
    the two rows share a hymnal number or a title but stay apart; 'same hymn, own
    rows': they are one hymn, but a version of it keeps a row of its own (the English
    Te Deum beside 'Te Deum laudamus'), so they too stay apart. `identified_by` (a
    hymnal number, 'TLH 252') says which row is meant when several bear
    `title`."""
    if not SAME_HYMN_CSV.exists():
        return []
    out = []
    for r in read_csv(SAME_HYMN_CSV):
        decision = r["decision"]
        if decision not in ("same hymn", "different hymns", "same hymn, own rows"):
            raise SystemExit(f"{SAME_HYMN_CSV.name}: decision must be 'same hymn', 'different hymns' or "
                             f"'same hymn, own rows', not {decision!r}")
        out.append((r["title"], r["master_title"], decision, (r.get("identified_by") or "").strip()))
    return out


def row_number_names(con):
    """{hymn_id: {'TLH 251', 'TL 157a', ...}}: each row's numbers as identified_by gives them."""
    out = defaultdict(set)
    for hymn_id, code, n, suffix in con.execute("SELECT hymn_id, hymnal_code, number, suffix FROM hymn_numbers"):
        out[hymn_id].add(f"{code} {n}{suffix}")
    return out


def titled_rows(con):
    """{title: {hymn_id}}: the rows with that HYMN TITLE, or if none, the rows
    with it as an Alternative English Title."""
    own, alternative = defaultdict(set), defaultdict(set)
    for hymn_id, title in con.execute("SELECT id, title FROM hymns"):
        own[title].add(hymn_id)
    for hymn_id, title in con.execute("SELECT hymn_id, title FROM hymn_alternative_titles"):
        alternative[title].add(hymn_id)
    return defaultdict(set, {t: set(own.get(t) or alternative[t]) for t in set(own) | set(alternative)})


def title_rows(con):
    """{title: {hymn_id}}: the rows a title in a decision or correction file
    means. That is the rows with that HYMN TITLE (or if none, with it as an
    Alternative English Title); the row another title was merged into
    (same_hymn_decisions.csv), also before write_back.py has merged it; and a
    row master_corrections.csv retitles, by its old title and by its new one."""
    out = titled_rows(con)
    numbers = row_number_names(con)
    merged = [(a, b, which) for a, b, decision, which in read_same_hymn_decisions() if decision == "same hymn"]
    renamed = [(r["old"], r["new"]) for r in read_csv(MASTER_CORRECTIONS_CSV) if r["hymnal"] == TITLE_CORRECTION]
    while True:
        before = sum(len(v) for v in out.values())
        for old, new in renamed:
            both = out[old] | out[new]
            out[old], out[new] = set(both), set(both)
        for a, b, which in merged:
            out[a] |= {h for h in out[b] if not which or which in numbers[h]}
        if sum(len(v) for v in out.values()) == before:
            return out


def read_link_decisions(con):
    """link_decisions.csv, a person's (or the research's) verdict on links the
    matching couldn't settle, as [(hymn_id, code, number, decision, reason)].
    `identified_by` (a hymnal number, 'ACS 1019') says which row is meant when
    several bear master_title."""
    rows_titled = title_rows(con)
    numbers = row_number_names(con)
    out = []
    for r in read_csv(LINK_DECISIONS_CSV):
        code, number, title, decision = r["hymnal"], parse_number(r["number"]), r["master_title"], r["decision"]
        if decision not in ("match", "no match"):
            raise SystemExit(f"link_decisions.csv: decision must be 'match' or 'no match', not {decision!r}")
        which = (r.get("identified_by") or "").strip()
        rows = {h for h in rows_titled.get(title, ()) if not which or which in numbers[h]}
        if not rows:  # e.g. a row write_back.py hasn't added yet
            print(f"{LINK_DECISIONS_CSV.name}: no master-index row is titled {title!r}"
                  + (f" with {which}" if which else "") + "; skipped")
        out += [(hymn_id, code, number, decision, r["reason"]) for hymn_id in sorted(rows)]
    return out


def same_hymn_groups(con):
    """The rows same_hymn_decisions.csv makes one hymn, as {hymn_id: group id}, and
    the pairs it keeps apart, as {frozenset((hymn_id, hymn_id))}. Several rows
    bearing a title are all meant unless identified_by picks one."""
    titled = titled_rows(con)
    numbers = row_number_names(con)
    parent = {}

    def find(x):
        parent.setdefault(x, x)
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    apart = set()
    for a, b, decision, which in read_same_hymn_decisions():
        rows_a = {h for h in titled.get(a, ()) if not which or which in numbers[h]}
        rows_b = titled.get(b, set())
        if decision == "same hymn":
            for y in rows_b:  # a group even once `a` is merged in (or spelled like `b`, so not kept)
                find(y)
        for x in rows_a:
            for y in rows_b:
                if decision == "same hymn":
                    parent[find(x)] = find(y)
                elif x != y:
                    apart.add(frozenset((x, y)))
    return {h: find(h) for h in parent}, apart


def report_shared_numbers(con):
    """Pairs of rows that share a hymnal number (page numbers aside) though
    same_hymn_decisions.csv neither merges them nor keeps them apart: the same
    hymn under two titles, or a number on the wrong row. As review rows."""
    group, apart = same_hymn_groups(con)
    shared = defaultdict(list)
    by_number = defaultdict(set)
    for hymn_id, code, n, suffix in con.execute(
        "SELECT hymn_id, hymnal_code, number, suffix FROM hymn_numbers WHERE is_page = 0"
    ):
        by_number[code, link_number(n, suffix) if code in LETTERED_HYMNALS else n].add(hymn_id)
    for (code, n), rows in by_number.items():
        rows = sorted(rows)
        for i, x in enumerate(rows):
            for y in rows[i + 1:]:
                pair = frozenset((x, y))
                if pair in apart or (x in group and y in group and group[x] == group[y]):
                    continue
                shared[x, y].append((code, n))
    titles = dict(con.execute("SELECT id, title FROM hymns"))
    order = {code: i for i, code in enumerate(HYMNALS[h][0] for h in HYMNALS)}
    out = []
    for (x, y), nums in sorted(shared.items()):
        nums.sort(key=lambda cn: (order[cn[0]], number_order(cn[1])))
        listed = ", ".join(f"{c} {n}" for c, n in nums)
        code, n = nums[0]
        out.append((code, n, titles[y], x, titles[x], "same hymn?",
                    f"'{titles[x]}' and '{titles[y]}' (row {y}) both have {listed}: one hymn under two titles "
                    f"(merge them in {SAME_HYMN_CSV.name}), different hymns ('different hymns' there), "
                    f"one hymn kept on two rows ('same hymn, own rows') "
                    f"or a number on the wrong row (master_corrections.csv)"))
    return out


def apply_link_decisions(con, decisions, add=True):
    """'match' makes the link clean, adding it (when `add`) if no route found
    it; 'no match' rejects it. A rejected number already in the master index
    is taken out of its cell by write_back.py. Applying them again is harmless.
    A decision on a hymnal the master index has no column for yet waits for
    write_back.py --add to put it there."""
    present = {c for (c,) in con.execute("SELECT code FROM hymnals")}
    for hymn_id, code, number, decision, reason in decisions:
        if code not in present:
            continue
        key = (hymn_id, code, number)
        status, note = ("ok" if decision == "match" else "rejected"), f"decided {decision}: {reason}"
        found = con.execute(
            "SELECT status, detail FROM links WHERE hymn_id = ? AND hymnal_code = ? AND number = ?", key
        ).fetchone()
        if found is None:
            if decision == "match" and add:
                con.execute("INSERT INTO links VALUES (?,?,?,?,?,?)", (*key, "decision", "ok", note))
        elif found[0] != status or note not in found[1]:
            con.execute(
                "UPDATE links SET status = ?, detail = detail || ? "
                "WHERE hymn_id = ? AND hymnal_code = ? AND number = ?",
                (status, "" if note in found[1] else f"; {note}", *key),
            )


def report_unlinked(con, review):
    """KELG and ELHB hymns left with no link once decisions are applied."""
    listed = {(code, n) for code, n, _, _ in review}
    for code, table, title in (("KELG", "kelg_hymns", "german_title"), ("ELHB", "elhb_hymns", "title")):
        linked = {n for (n,) in con.execute(
            "SELECT number FROM links WHERE hymnal_code = ? AND status != 'rejected'", (code,))}
        for n, t in con.execute(f"SELECT number, {title} FROM {table} ORDER BY number").fetchall():
            if n not in linked and (code, n) not in listed:
                review.append((code, n, t, "no master-index row found"))


def sort_tlhy(con):
    """Fill tlhy_sorting: each TLHy hymn as a KELG hymn (its rows carry a KELG
    number), a translation from another language (tlhy_originals.csv), an
    English original (no translator credited, or researched as English), or
    a translation whose original wasn't traced."""
    links = defaultdict(lambda: defaultdict(set))  # number -> status -> rows
    for hymn_id, number, status in con.execute(
        "SELECT hymn_id, number, status FROM links WHERE hymnal_code = 'TLHy'"
    ):
        links[number][status].add(hymn_id)
    row_kelg = defaultdict(set)
    for hymn_id, n in con.execute(
        "SELECT hymn_id, number FROM hymn_numbers WHERE hymnal_code = 'KELG' "
        "UNION SELECT hymn_id, number FROM links WHERE hymnal_code = 'KELG' AND status = 'ok'"
    ):
        row_kelg[hymn_id].add(n)
    originals = {r[0]: r[1:] for r in con.execute(
        "SELECT number, original_title, language, author FROM tlhy_originals")}
    for number, first_line, translator in con.execute(
        "SELECT number, first_line, translator FROM tlhy_index WHERE alternate = 0"
    ).fetchall():
        ok = links[number]["ok"]
        status = "matched" if ok else ("needs review" if links[number]["review"] else "unmatched")
        kelg = set().union(*(row_kelg[h] for h in ok)) if ok else set()
        title, language, author = originals.get(number, (None, None, None))
        if kelg:
            category = "KELG hymn"
        elif language and language != "English":
            category = f"translated from {language}"
        elif not translator or language == "English":
            category = "English original"
        else:
            category = "translation, original not traced"
        con.execute("INSERT INTO tlhy_sorting VALUES (?,?,?,?,?,?,?,?,?)",
                    (number, first_line, translator, status, category,
                     fmt_nums(kelg) if kelg else None, title, language, author))


def flag_master_conflicts(con):
    """A proposed link is only clean if the row has no other number in that column."""
    master = defaultdict(set)
    for hymn_id, code, number, suffix in con.execute(
        "SELECT hymn_id, hymnal_code, number, suffix FROM hymn_numbers WHERE hymnal_code IN ({})".format(
            ", ".join(f"'{c}'" for c in LINKED_HYMNALS))
    ):
        master[hymn_id, code].add(link_number(number, suffix) if code in LETTERED_HYMNALS else number)
    for hymn_id, code, number in con.execute(
        "SELECT hymn_id, hymnal_code, number FROM links WHERE method != 'in master'"
    ).fetchall():
        have = master.get((hymn_id, code))
        if have and number not in have:
            con.execute(
                "UPDATE links SET status='review', detail=detail || ? "
                "WHERE hymn_id=? AND hymnal_code=? AND number=?",
                (f"; master index already has {code} {fmt_nums(have)}", hymn_id, code, number),
            )


# --- main --------------------------------------------------------------------

def build():
    if DB_PATH.exists():
        DB_PATH.unlink()
    con = sqlite3.connect(DB_PATH)
    con.executescript(SCHEMA)
    load_lcms(con)
    load_kelg(con)
    load_elhb(con)
    load_wh(con)
    load_tlhy(con)
    load_tlh(con)
    load_tlhy_originals(con)
    load_tl(con)
    load_tl_originals(con)
    load_elgbs_lpc(con)
    load_hymnary(con)
    load_kelg_english(con)
    fix_kelg_elhb_refs(con)
    build_hymn_titles(con)

    # The rows the master index has each number on. Such a number links to no
    # other row (but by a decision): the rows a hymnal's index matched are
    # settled, and a title it shares with a row added since (from a Hymnary.org
    # hymnal) doesn't make that row the same hymn.
    placed_rows = defaultdict(set)
    for hymn_id, code, n, suffix in con.execute(
        "SELECT hymn_id, hymnal_code, number, suffix FROM hymn_numbers WHERE is_page = 0"
    ):
        placed_rows[code, link_number(n, suffix) if code in LETTERED_HYMNALS else n].add(hymn_id)

    def add_link(hymn_id, code, number, method, status, detail):
        rows = placed_rows.get((code, number))
        if rows and hymn_id not in rows:
            return
        # First route to find a link wins; a later 'review' still flags it, and a
        # later clean route clears a flag on a number already in the master index.
        existing = con.execute(
            "SELECT status, method FROM links WHERE hymn_id=? AND hymnal_code=? AND number=?",
            (hymn_id, code, number),
        ).fetchone()
        if existing is None:
            con.execute("INSERT INTO links VALUES (?,?,?,?,?,?)",
                        (hymn_id, code, number, method, status, detail))
        elif existing == ("review", "in master") and status == "ok":
            con.execute(
                "UPDATE links SET status='ok', detail=detail || '; confirmed ' || ? "
                "WHERE hymn_id=? AND hymnal_code=? AND number=?",
                (f"{method} ({detail})", hymn_id, code, number),
            )
        elif status == "review" and existing[0] == "ok":
            con.execute(
                "UPDATE links SET status='review', detail=detail || '; ' || ? "
                "WHERE hymn_id=? AND hymnal_code=? AND number=?",
                (detail, hymn_id, code, number),
            )

    review = []  # (source, number, title, problem): hymns that couldn't be placed
    decisions = read_link_decisions(con)
    placed = link_from_master(con, add_link)
    kelg_links = link_kelg(con, add_link, review)
    elhb_found = link_elhb(con, kelg_links, placed, add_link, review)
    link_kelg_via_elhb(con, elhb_found, add_link)
    link_kelg_by_wh(con, kelg_links, add_link)
    apply_link_decisions(con, decisions, add=False)  # the TLHy and TL routes build on the KELG links
    link_tlhy(con, add_link)
    confirm_tlhy_by_translator(con)
    link_by_original(con, add_link, "TLHy", "tlhy_originals")
    link_tl(con, add_link)
    link_by_original(con, add_link, "TL", "tl_originals")
    link_lpc(con, add_link)
    link_hymnary(con, add_link, decisions)
    if "LH" in hymnary_codes(con):
        link_by_original(con, add_link, "LH", "lh_originals")
    flag_master_conflicts(con)
    apply_link_decisions(con, decisions)
    report_unmatched_tlhy(con, review)
    report_unmatched_tl(con, review)
    report_unmatched_lpc(con, review)
    report_unmatched_hymnary(con, review)
    report_unlinked(con, review)
    sort_tlhy(con)

    con.execute(
        """CREATE VIEW kelg_index AS
SELECT k.number, k.german_title, k.wh_title, k.verses, k.supplement,
    MAX(CASE WHEN r.hymnal_code = 'ELHB' THEN COALESCE(r.fixed, r.raw) END) AS ELHB,
    MAX(CASE WHEN r.hymnal_code = 'ELHB' THEN r.title END) AS ELHB_title,
    MAX(CASE WHEN r.hymnal_code = 'TLH' THEN r.raw END) AS TLH,
    MAX(CASE WHEN r.hymnal_code = 'TLH' THEN r.title END) AS TLH_title,
    MAX(CASE WHEN r.hymnal_code = 'LW' THEN r.raw END) AS LW,
    MAX(CASE WHEN r.hymnal_code = 'LW' THEN r.title END) AS LW_title,
    MAX(CASE WHEN r.hymnal_code = 'ELH' THEN r.raw END) AS ELH,
    MAX(CASE WHEN r.hymnal_code = 'ELH' THEN r.title END) AS ELH_title,
    MAX(CASE WHEN r.hymnal_code = 'LSB' THEN r.raw END) AS LSB,
    MAX(CASE WHEN r.hymnal_code = 'LSB' THEN r.title END) AS LSB_title,
    (SELECT group_concat(e.title, ' / ') FROM kelg_english e WHERE e.kelg_number = k.number)
        AS other_english_title,
    (SELECT group_concat(COALESCE(e.translator || ' (' || e.source || ')', e.source), ' / ')
     FROM kelg_english e WHERE e.kelg_number = k.number) AS other_english_source
FROM kelg_hymns k LEFT JOIN kelg_refs r ON r.kelg_number = k.number
GROUP BY k.number
ORDER BY k.number"""
    )
    con.execute(
        """CREATE VIEW tlhy_status AS
SELECT t.number, t.first_line, t.translator,
    CASE WHEN EXISTS (SELECT 1 FROM links l WHERE l.hymnal_code = 'TLHy' AND l.number = t.number
                      AND l.status = 'ok') THEN 'matched'
         WHEN EXISTS (SELECT 1 FROM links l WHERE l.hymnal_code = 'TLHy' AND l.number = t.number
                      AND l.status = 'review') THEN 'needs review'
         ELSE 'unmatched' END AS status,
    (SELECT group_concat(h.title || ' [' || l.status || ']', ' | ')
     FROM links l JOIN hymns h ON h.id = l.hymn_id
     WHERE l.hymnal_code = 'TLHy' AND l.number = t.number AND l.status != 'rejected') AS master_rows
FROM tlhy_index t WHERE t.alternate = 0
ORDER BY t.number"""
    )
    con.execute(
        """CREATE VIEW tl_status AS
SELECT t.number, t.title,
    CASE WHEN EXISTS (SELECT 1 FROM links l WHERE l.hymnal_code = 'TL' AND l.number = t.number
                      AND l.status = 'ok') THEN 'matched'
         WHEN EXISTS (SELECT 1 FROM links l WHERE l.hymnal_code = 'TL' AND l.number = t.number
                      AND l.status = 'review') THEN 'needs review'
         ELSE 'unmatched' END AS status,
    (SELECT group_concat(h.title || ' [' || l.status || ']', ' | ')
     FROM links l JOIN hymns h ON h.id = l.hymn_id
     WHERE l.hymnal_code = 'TL' AND l.number = t.number AND l.status != 'rejected') AS master_rows
FROM tl_index t WHERE t.alternate = 0
ORDER BY CAST(t.number AS INTEGER), t.number"""
    )
    con.execute(
        """CREATE VIEW hymnary_status AS
SELECT w.hymnal, w.number, w.first_line, w.english, w.original, w.tune,
    CASE WHEN EXISTS (SELECT 1 FROM links l WHERE l.hymnal_code = w.hymnal AND l.number = w.number
                      AND l.status = 'ok') THEN 'matched'
         WHEN EXISTS (SELECT 1 FROM links l WHERE l.hymnal_code = w.hymnal AND l.number = w.number
                      AND l.status = 'review') THEN 'needs review'
         ELSE 'unmatched' END AS status,
    (SELECT group_concat(h.title || ' [' || l.status || ']', ' | ')
     FROM links l JOIN hymns h ON h.id = l.hymn_id
     WHERE l.hymnal_code = w.hymnal AND l.number = w.number AND l.status != 'rejected') AS master_rows
FROM hymnary_index w WHERE w.alternate = 0
ORDER BY w.rowid"""
    )
    hymnary_proposed = "".join(
        f"""    (SELECT group_concat(number, ', ') FROM
        (SELECT number FROM links WHERE hymn_id = i.id AND hymnal_code = '{code}' AND status != 'rejected'
         ORDER BY CAST(number AS INTEGER), number))
        AS proposed_{code},
""" for code in HYMNARY_CSVS)
    con.execute(
        f"""CREATE VIEW proposed_index AS
SELECT i.*,
    (SELECT group_concat(number, ', ') FROM
        (SELECT number FROM links WHERE hymn_id = i.id AND hymnal_code = 'KELG' AND status != 'rejected'
         ORDER BY number))
        AS proposed_KELG,
    (SELECT group_concat(number, ', ') FROM
        (SELECT number FROM links WHERE hymn_id = i.id AND hymnal_code = 'ELHB' AND status != 'rejected'
         ORDER BY number))
        AS proposed_ELHB,
    (SELECT group_concat(number, ', ') FROM
        (SELECT number FROM links WHERE hymn_id = i.id AND hymnal_code = 'TLHy' AND status != 'rejected'
         ORDER BY number))
        AS proposed_TLHy,
    (SELECT group_concat(number, ', ') FROM
        (SELECT number FROM links WHERE hymn_id = i.id AND hymnal_code = 'TL' AND status != 'rejected'
         ORDER BY CAST(number AS INTEGER), number))
        AS proposed_TL,
    (SELECT group_concat(number, ', ') FROM
        (SELECT number FROM links WHERE hymn_id = i.id AND hymnal_code = 'LPC' AND status != 'rejected'
         ORDER BY number))
        AS proposed_LPC,
{hymnary_proposed}    (SELECT group_concat(german_title, ' / ') FROM
        (SELECT DISTINCT k.german_title FROM links l JOIN kelg_hymns k ON k.number = l.number
         WHERE l.hymn_id = i.id AND l.hymnal_code = 'KELG' AND l.status != 'rejected' ORDER BY k.number))
        AS proposed_german_title,
    (SELECT group_concat(english_first_line, ' / ') FROM
        (SELECT DISTINCT w.english_first_line FROM links l JOIN wh_hymns w ON w.number = l.number
         WHERE l.hymn_id = i.id AND l.hymnal_code = 'KELG' AND l.status = 'ok' ORDER BY w.number))
        AS wh_english_first_line,
    (SELECT COUNT(*) FROM links WHERE hymn_id = i.id AND status = 'review') AS needs_review
FROM hymn_index i"""
    )
    con.commit()
    shared = report_shared_numbers(con)
    write_review(con, review, shared)
    report(con, review, shared)
    con.close()


def write_review(con, review, shared=()):
    rows = [
        (code, number, source_title(con, code, number), hymn_id, title, method, detail)
        for code, number, hymn_id, title, method, detail in con.execute(
            """SELECT l.hymnal_code, l.number, l.hymn_id, h.title, l.method, l.detail
               FROM links l JOIN hymns h ON h.id = l.hymn_id
               WHERE l.status = 'review' ORDER BY l.hymnal_code DESC, l.number, l.hymn_id"""
        ).fetchall()
    ]
    rows += [(code, number, title, None, None, "unmatched", problem)
             for code, number, title, problem in review]
    rows += list(shared)
    with open(REVIEW_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["source", "number", "source_title", "master_row", "master_title", "method", "detail"])
        w.writerows(rows)


def source_title(con, code, number):
    if code in HYMNARY_CSVS:
        row = con.execute("SELECT first_line FROM hymnary_index WHERE hymnal = ? AND number = ? "
                          "ORDER BY alternate", (code, number)).fetchone()
        return row[0] if row else None
    sql = {
        "KELG": "SELECT german_title FROM kelg_hymns WHERE number=?",
        "ELHB": "SELECT title FROM elhb_hymns WHERE number=?",
        "TLHy": "SELECT first_line FROM tlhy_index WHERE number=? ORDER BY alternate",
        "TL": "SELECT title FROM tl_index WHERE number=? ORDER BY alternate",
        "LPC": "SELECT english_first_line FROM elgbs_lpc_hymns WHERE number=?",
    }[code]
    row = con.execute(sql, (number,)).fetchone()
    return row[0] if row else None


def report(con, review, shared=()):
    q = lambda sql: con.execute(sql).fetchone()[0]
    codes = hymnary_codes(con)
    hymnary = {code: q(f"SELECT COUNT(DISTINCT number) FROM hymnary_index WHERE hymnal = '{code}'")
               for code in codes}
    print(f"wrote {DB_PATH.name}: {q('SELECT COUNT(*) FROM hymns')} master rows, "
          f"{q('SELECT COUNT(*) FROM kelg_hymns')} KELG hymns, {q('SELECT COUNT(*) FROM elhb_hymns')} ELHB hymns, "
          f"{q('SELECT COUNT(DISTINCT number) FROM tlhy_index')} TLHy hymns, "
          f"{q('SELECT COUNT(DISTINCT number) FROM tl_index')} TL hymns, "
          f"{q('SELECT COUNT(*) FROM elgbs_lpc_hymns')} LPC/ELGBS hymns, "
          + ", ".join(f"{n} {code}" for code, n in hymnary.items()) + " hymns; "
          f"{q('SELECT COUNT(*) FROM hymn_alternative_titles')} alternative titles")
    for code, table in (("KELG", "kelg_hymns"), ("ELHB", "elhb_hymns"), ("TLHy", "tlhy_index"),
                        ("TL", "tl_index"), ("LPC", "elgbs_lpc_hymns"),
                        *((code, f"hymnary_index WHERE hymnal = '{code}'") for code in codes)):
        total = q(f"SELECT COUNT(DISTINCT number) FROM {table}")
        ok = q(f"SELECT COUNT(DISTINCT number) FROM links WHERE hymnal_code='{code}' AND status='ok'")
        any_ = q(f"SELECT COUNT(DISTINCT number) FROM links WHERE hymnal_code='{code}' AND status != 'rejected'")
        rows = q(f"SELECT COUNT(DISTINCT hymn_id) FROM links WHERE hymnal_code='{code}' AND status != 'rejected'")
        print(f"  {code}: {any_}/{total} hymns linked to {rows} master rows "
              f"({ok} with at least one clean link)")
    sorting = con.execute(
        "SELECT category, COUNT(*) FROM tlhy_sorting GROUP BY category ORDER BY COUNT(*) DESC").fetchall()
    if sorting:
        print("  TLHy sorted: " + ", ".join(f"{n} {c}" for c, n in sorting))
    to_review = q("SELECT COUNT(*) FROM links WHERE status='review'")
    rejected = q("SELECT COUNT(*) FROM links WHERE status='rejected'")
    print(f"  {to_review} links + {len(review)} other items + {len(shared)} pairs of rows sharing a number "
          f"to review -> {REVIEW_CSV.name}; {rejected} links rejected in {LINK_DECISIONS_CSV.name}")


if __name__ == "__main__":
    build()
