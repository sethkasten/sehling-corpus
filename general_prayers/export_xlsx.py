# -*- coding: utf-8 -*-
"""Export general_prayers.db to an Excel workbook for comparing the general
prayers of the church orders petition by petition."""
import sqlite3, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
DB  = os.path.join(HERE, '..', 'general_prayers.db')
OUT = os.path.join(HERE, '..', 'general_prayers.xlsx')

FONT = 'Arial'
HDR_FILL = PatternFill('solid', fgColor='1F3864')
HDR_FONT = Font(name=FONT, size=10, bold=True, color='FFFFFF')
BODY     = Font(name=FONT, size=10)
BOLD     = Font(name=FONT, size=10, bold=True)
TITLE    = Font(name=FONT, size=14, bold=True, color='1F3864')
ALT_FILL = PatternFill('solid', fgColor='F2F5FA')
HIT_FILL = PatternFill('solid', fgColor='C9D7EE')    # primary category present
SUB_FILL = PatternFill('solid', fgColor='E8EEF8')    # secondary only
THIN     = Side(style='thin', color='BFBFBF')


def style_header(ws, ncols, row=1, freeze='B2'):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = freeze


def table(ws, heads, rows, widths, wrap=(), freeze='B2', band_key=None):
    ws.append(heads)
    band, prev = False, object()
    for r in rows:
        ws.append(list(r))
        k = band_key(r) if band_key else None
        if k != prev: band, prev = not band, k
        for i, c in enumerate(ws[ws.max_row], 1):
            c.font = BODY
            c.alignment = Alignment(vertical='top', wrap_text=i in wrap)
            if band_key and band: c.fill = ALT_FILL
    style_header(ws, len(heads), freeze=freeze)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f'A1:{get_column_letter(len(heads))}{ws.max_row}'


def main():
    db = sqlite3.connect(DB)
    q = lambda sql, *a: db.execute(sql, a).fetchall()
    nw, = q('SELECT COUNT(*) FROM witnesses')[0]
    np_, = q('SELECT COUNT(*) FROM petitions')[0]
    nf, = q('SELECT COUNT(*) FROM families')[0]
    y0, y1 = q('SELECT MIN(year), MAX(year) FROM witnesses')[0]
    wb = Workbook()

    # ---------------- Read me ----------------
    ws = wb.active; ws.title = 'Read me'
    notes = [
        ('Sehling Corpus — the general prayer (Prayers of the Church) of the Sunday service', TITLE),
        ('', BODY),
        (f'{np_} petitions · {nw} witnesses (church orders) · {nf} text families · {y0}–{y1}', BOLD),
        ('Generated from general_prayers.db. Rebuild with general_prayers/build_gp_db.py, then general_prayers/export_xlsx.py.', BODY),
        ('', BODY),
        ('Sheets', BOLD),
        ('Compare — one row per order, one column per standard category; each cell gives the petition number(s) at which', BODY),
        ('      that intention occurs: 3 = the petition\'s chief intention, (3) = named in passing within petition 3. Start here.', BODY),
        ('Prayers — every petition: prayer text (original + formal-equivalence English), bid/rubric (original + English),', BODY),
        ('      source citation and standard category. Filter on Category or Witness key.', BODY),
        ('Witnesses — each order: date, territory, citation, family, form, liturgical position, the sequence of categories.', BODY),
        ('Families — the text families (who copied whom), with the archetype of each.', BODY),
        ('Categories — the standard category scheme, with how often each is used.', BODY),
        ('', BODY),
        ('How the texts were made', BOLD),
        ('Original-language texts are the base text of Sehling\'s edition with his apparatus (sigla, variant editions,', BODY),
        ('      footnotes, page/folio markers) removed; every text was machine-checked against the corpus (≥ 90 % token match).', BODY),
        ('English is a deliberately literal rendering in the idiom of the Authorized Version (thee/thou, -eth), keeping', BODY),
        ('      word order and clause structure where English allows. Where one order copies another, the translation is', BODY),
        ('      reused and patched only where the German differs (column "Translation reused from").', BODY),
        ('"Bid" = the bidding addressed to the people ("Let us pray for …"); "rubric" = the heading or marginal note.', BODY),
        ('', BODY),
        ('Things that will bite you', BOLD),
        ('Dates are those of the order as Sehling edits it; several corpus documents carry a misleading year (see Witnesses notes).', BODY),
        ('Collect banks, the Litany, prayer-day and occasional prayers, and weekday/afternoon forms are excluded.', BODY),
        ('Petition boundaries follow the witness: a continuous prayer is split at its own "Item / Auch / Und" joints.', BODY),
        ('', BODY),
        ('See GENERAL_PRAYERS_GUIDE.md in the repository for the full account.', BODY),
    ]
    for i, (text, font) in enumerate(notes, start=1):
        ws.cell(row=i, column=1, value=text).font = font
    ws.column_dimensions['A'].width = 125

    # ---------------- Compare ----------------
    ws = wb.create_sheet('Compare')
    cats = q('SELECT code, label FROM categories ORDER BY sort')
    heads = ['Order', 'Year', 'Family', 'Petitions'] + [lab for _, lab in cats]
    ws.append(heads)
    rows = q('SELECT w.territory, w.year, w.family_code, w.n_petitions, '
             + ', '.join(f'cp."{c}"' for c, _ in cats) +
             ', w.key FROM category_pivot cp JOIN witnesses w ON w.key = cp.witness_key '
             'ORDER BY CASE WHEN w.family_code LIKE \'R%\' THEN 1 ELSE 0 END, w.family_code, w.year, w.key')
    band, prev = False, None
    for r in rows:
        key = r[-1]
        vals = list(r[:-1])
        vals[0] = f'{vals[0]} ({key})'
        ws.append(vals)
        if r[2] != prev: band, prev = not band, r[2]
        for i, c in enumerate(ws[ws.max_row], 1):
            c.font = BODY
            c.alignment = Alignment(vertical='top', horizontal='left' if i <= 3 else 'center')
            c.border = Border(bottom=THIN)
            if i > 4 and c.value:
                c.fill = HIT_FILL if any(t[0] != '(' for t in str(c.value).split()) else SUB_FILL
            elif band and i <= 4:
                c.fill = ALT_FILL
    # usage row
    ws.append(['Orders with this intention', None, None, None] +
              [q('SELECT witnesses_any FROM category_usage WHERE code=?', c)[0][0] for c, _ in cats])
    for c in ws[ws.max_row]: c.font = BOLD; c.alignment = Alignment(horizontal='center')
    ws[f'A{ws.max_row}'].alignment = Alignment(horizontal='left')
    style_header(ws, len(heads), freeze='E2')
    ws.row_dimensions[1].height = 75
    for c in ws[1][4:]:
        c.alignment = Alignment(text_rotation=90, vertical='bottom', horizontal='center', wrap_text=True)
    for i, w in enumerate([58, 7, 7, 9] + [6] * len(cats), 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ---------------- Prayers ----------------
    ws = wb.create_sheet('Prayers')
    heads = ['Year', 'Order', 'Territory', 'No.', 'Category', 'Also names',
             'Prayer (original)', 'Prayer (English)', 'Source',
             'Bid / rubric (original)', 'Bid / rubric (English)',
             'Family', 'Form', 'Tradition', 'Note', 'Translation reused from', 'Verified', 'Witness key']
    lab = dict(cats)
    rows = []
    for r in q('SELECT w.year, w.order_title, w.territory, p.seq, c.label, p.subcategories, '
               'p.prayer_original, p.prayer_english, w.citation, '
               'v.bid_rubric_original, v.bid_rubric_english, '
               'w.family, w.form, w.tradition, p.note, p.translation_reused_from, p.verify_coverage, w.key '
               'FROM petitions p JOIN witnesses w ON w.key = p.witness_key '
               'JOIN categories c ON c.code = p.category '
               'JOIN prayers_of_the_church v ON v.witness_key = w.key AND v.seq = p.seq '
               'ORDER BY w.year, w.key, p.seq'):
        r = list(r)
        if r[5]: r[5] = '; '.join(lab[s] for s in r[5].split('; '))
        for i in (9, 10):
            r[i] = r[i] or None
        rows.append(r)
    table(ws, heads, rows,
          [7, 40, 26, 5, 22, 22, 70, 70, 40, 45, 45, 30, 26, 14, 40, 20, 8, 24],
          wrap={2, 3, 5, 6, 7, 8, 9, 10, 11, 15}, freeze='E2', band_key=lambda r: r[-1])

    # ---------------- Witnesses ----------------
    ws = wb.create_sheet('Witnesses')
    heads = ['Year', 'Order', 'Territory', 'Family', 'Tradition', 'Form', 'Petitions',
             'Categories in sequence', 'Position in the service', 'Heading (original)',
             'Heading (English)', 'Citation', 'Notes', 'eko.db doc', 'Witness key']
    rows = q('SELECT year, order_title, territory, family, tradition, form, n_petitions, '
             'categories_sequence, position, heading_original, heading_english, citation, notes, '
             'eko_doc_id, key FROM witnesses ORDER BY year, key')
    table(ws, heads, rows, [7, 40, 28, 30, 14, 26, 9, 60, 60, 40, 40, 40, 70, 10, 24],
          wrap={2, 3, 4, 6, 8, 9, 10, 11, 12, 13}, freeze='C2')

    # ---------------- Families ----------------
    ws = wb.create_sheet('Families')
    heads = ['Code', 'Family', 'Archetype', 'Witnesses', 'Description', 'Members']
    rows = [list(r) + ['; '.join(f'{t} {y}' for t, y in q(
                'SELECT territory, year FROM witnesses WHERE family_code=? ORDER BY year', r[0]))]
            for r in q('SELECT code, name, archetype, n_witnesses, description FROM families '
                       'ORDER BY CASE WHEN code LIKE \'R%\' THEN 1 ELSE 0 END, code')]
    table(ws, heads, rows, [6, 36, 30, 10, 80, 70], wrap={2, 3, 5, 6})

    # ---------------- Categories ----------------
    ws = wb.create_sheet('Categories')
    heads = ['Code', 'Label', 'Description', 'Petitions (chief intention)', 'Orders naming it']
    rows = q('SELECT u.code, u.label, c.description, u.petitions_primary, u.witnesses_any '
             'FROM category_usage u JOIN categories c ON c.code = u.code ORDER BY u.sort')
    table(ws, heads, rows, [16, 28, 90, 14, 14], wrap={3})

    wb.save(OUT)
    print(f'{OUT}: {np_} petitions, {nw} witnesses')


if __name__ == '__main__':
    main()
